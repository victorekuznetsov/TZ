# -*- coding: utf-8 -*-
"""Фикстура tests/fixtures/pm06_sample.xlsx — выборка из настоящей выгрузки
PM-06 (M06_1200_2027.xlsx ветки rawdata TOPO) для tests/test_toro_rebuild.mjs.

В выборку попадают ВСЕ строки выбранных заказов (иначе суммы заказа не
совпадут с витринами): заказы WK Сухого Лога на разных стадиях — позиция ППР
без заказа, на согласовании, согласован (СГГС), оригинал БЕ с копией
«Развития» и сама копия, заказы с признаком ППМ «Начиная с деблок.» /
«Никогда», — и немного строк балансовой единицы 1200 (разводка по БЕ).
Ожидаемые значения теста — из витрин, собранных Python.

    python3 tests/make_pm06_fixture.py <M06_1200_2027.xlsx>
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PICK = ["110000023773", "1102763037", "1102763038", "1102277367", "1102688138", "1200434254"]


def main():
    src = sys.argv[1]
    prov = json.load(open(os.path.join(ROOT, "data", "provision.json"), encoding="utf-8"))
    ctl = json.load(open(os.path.join(ROOT, "data", "control.json"), encoding="utf-8"))
    ix = {c: i for i, c in enumerate(ctl["columns"])}
    D = ctl["dictionaries"]
    g = lambda r, c: D[c][r[ix[c]]] if c in D else r[ix[c]]
    wk2400 = {g(r, "order") for r in ctl["rows"] if g(r, "y") == "2027" and g(r, "plant") == "2400"}
    flagged = []
    for k, v in prov["meta"]["ppmFlags"].items():
        y, o, _ = k.split("|")
        if y == "2027" and o in wk2400 and o not in flagged:
            flagged.append(o)
    pick = set(PICK) | set(flagged[:4])

    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb["PM-06"]
    head, body, hdr = [], [], None
    rows = ws.iter_rows(values_only=True)
    for r in rows:
        head.append(r)
        if any(isinstance(v, str) and v.strip() == "Заказ" for v in r):
            hdr = {v.strip(): i for i, v in enumerate(r) if isinstance(v, str) and v.strip()}
            break
    co, cbe = hdr["Заказ"], hdr["Балансовая единица"]
    ceo, cw, cbs = hdr["ЕО"] + 1, hdr["Заказ Вид работы ТОРО"], hdr["Заказ Базисный срок начала (дата)"]
    all_rows = list(rows)
    # пара копии: заказ «Развития» на ту же ЕО, вид работ и базисную дату, что оригинал
    keys = {(r[ceo], r[cw], r[cbs]) for r in all_rows if str(r[co]) in pick}
    n1200 = 0
    for r in all_rows:
        o = str(r[co])
        if o in pick or (r[ceo], r[cw], r[cbs]) in keys:
            body.append(r)
        elif r[cbe] == 'АО "Полюс Вернинское"' and n1200 < 40:
            body.append(r)
            n1200 += 1
    out = openpyxl.Workbook()
    sh = out.active
    sh.title = "PM-06"
    for r in head + body:
        sh.append(list(r))
    path = os.path.join(ROOT, "tests", "fixtures", "pm06_sample.xlsx")
    out.save(path)
    print(f"{path}: строк {len(body)}, заказов {len({str(r[co]) for r in body})}")


if __name__ == "__main__":
    main()
