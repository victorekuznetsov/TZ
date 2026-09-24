# -*- coding: utf-8 -*-
"""Презентация «Экскаваторы WK: исполнение 2024–2026 и план 2027» на шаблоне АО «Развитие».

Данные — из build_wk_status_data.py (витрины интерактивного отчёта WK CRM).
Запуск:
  python3 build/build_wk_status_data.py <TOPO/pm06_meta/order_status> wk.json
  python3 build/build_wk_status_deck.py <Развитие шаблон.pptx> wk.json <выход.pptx>
"""
import json
import re
import sys

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION

import deck_kit
from deck_kit import *  # noqa: F401,F403 — палитра, макеты и помощники вёрстки
from deck_kit import _nobullet

SRC, DATA, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
d = json.load(open(DATA, encoding='utf-8'))
prs = deck_kit.open_template(SRC)

# ── палитра графиков (проверена validate_palette.js: CVD и контраст; подписи на всех графиках)
C_STOCK, C_BUY, C_LATE, C_GAP = '12A06E', '5566C9', 'E0A030', 'D2454F'
C_PLAN, C_FACT = 'C3C8CE', '12A06E'
C_Y = {'2024': 'B9C1EE', '2025': '8793E0', '2026': '5566C9', '2027': '12A06E'}
C_GROUP = {'Закрыт': '0E7A54', 'Тех. закрыт': '12A06E', 'В работе': '7FD8B4', 'Деблокирован, пусто': C_LATE,
           'Согласование': C_BUY, 'ППР без заказа': 'C3C8CE', 'Нет статуса': 'E5E7EA'}
TINT_G, TINT_A, TINT_R, ZEBRA = 'D5F5E6', 'FCEFD2', 'F8DADC', 'F4F5F6'

SN = d['siteNames']
SHORT = {'1100': 'Красноярск', '1400': 'Магадан', '2400': 'Сухой Лог'}
SITES = list(SN)
YEARS = d['years']
FACT_YEARS = ['2024', '2025', '2026']
AS_OF = '14.09.2026'
M = 1e6


# ───────────────────────── форматирование чисел
def sp(n):
    return f'{n:,.0f}'.replace(',', ' ')


def mln(v, dec=0):
    if v is None:
        return '—'
    s = f'{v / M:,.{dec}f}'.replace(',', ' ').replace('.', ',')
    return s


def bln(v):
    return f'{v / 1e9:.2f}'.replace('.', ',')


def pct(v, dec=0):
    return '—' if v is None else f'{v * 100:.{dec}f}%'.replace('.', ',')


def unit_key(u):
    n = re.findall(r'\d+', u.split('№')[-1])
    return (u.split(' ')[0], int(n[0]) if n else 0)


FLEET = {f['unit']: f for f in d['fleet']}
UNITS = sorted(set(d['unit']) | set(FLEET), key=unit_key)


def site_of(u):
    return FLEET[u]['site'] if u in FLEET else '2400'


def model_of(u):
    return FLEET[u]['model'] if u in FLEET else u.split(' ')[0]


def units_of(site):
    return sorted([u for u in UNITS if site_of(u) == site], key=lambda u: (model_of(u), unit_key(u)[1]))


def U(u, y):
    return d['unit'].get(u, {}).get(y) or {'plan': 0, 'fact': 0, 'exec': None, 'noOrderPlan': 0, 'orders': 0}


def ktg(u, y, k):
    v = d['ktg']['unit'].get(u, {}).get(y, {}).get(k)
    # ровно 100% за весь год — заглушка витрины для бортов без данных
    return None if v is None or v >= 0.99999 else v


def prov_u(u, y):
    return d['prov']['unit'].get(u, {}).get(y)


def uncovered(p):
    return p['late'] + p['undated'] + p['gap']


# ───────────────────────── графики
def _font(obj, size=10, color=DARK, bold=False):
    obj.font.size = Pt(size)
    obj.font.name = FONT
    obj.font.bold = bold
    obj.font.color.rgb = rgb(color)


def _hide_labels(ser, idxs):
    dl = ser._element.get_or_add_dLbls()
    for i in sorted(idxs, reverse=True):
        e = etree.SubElement(dl, qn('c:dLbl'))
        etree.SubElement(e, qn('c:idx')).set('val', str(i))
        etree.SubElement(e, qn('c:delete')).set('val', '1')
        dl.remove(e)
        dl.insert(0, e)


def _tick_skip(chart, n):
    ax = chart._chartSpace.chart.plotArea.catAx_lst[0]
    nm = ax.find(qn('c:noMultiLvlLbl'))
    for tag in ('c:tickLblSkip', 'c:tickMarkSkip'):
        e = etree.Element(qn(tag))
        e.set('val', str(n))
        if nm is not None:
            nm.addprevious(e)
        else:
            ax.append(e)


def chart(s, x, y, w, h, kind, cats, series, fmt='# ##0', labels=True, legend=True, size=10, gap=60,
          overlap=None, vmax=None, vmin=None, label_pos=None, label_colors=None, min_label=None,
          value_axis=False, grid=False, axis_fmt=None, legend_pos='b', cat_size=None, tick_skip=None,
          line_width=2.25, point_colors=None, min_share=None):
    """series: [(имя, значения, цвет)]. Возвращает chart."""
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals, _ in series:
        cd.add_series(name, vals)
    t = {'col': XL_CHART_TYPE.COLUMN_CLUSTERED, 'colS': XL_CHART_TYPE.COLUMN_STACKED,
         'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'barS': XL_CHART_TYPE.BAR_STACKED,
         'bar100': XL_CHART_TYPE.BAR_STACKED_100, 'col100': XL_CHART_TYPE.COLUMN_STACKED_100,
         'line': XL_CHART_TYPE.LINE}[kind]
    gf = s.shapes.add_chart(t, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(size)
    ch.font.name = FONT
    ch.font.color.rgb = rgb(DARK)
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = {'b': XL_LEGEND_POSITION.BOTTOM, 't': XL_LEGEND_POSITION.TOP,
                              'r': XL_LEGEND_POSITION.RIGHT}[legend_pos]
        ch.legend.include_in_layout = False
        _font(ch.legend, size)
    pl = ch.plots[0]
    if kind != 'line':
        pl.gap_width = gap
        if overlap is not None:
            pl.overlap = overlap
        elif kind in ('colS', 'barS', 'bar100', 'col100'):
            pl.overlap = 100
    stacked = kind in ('colS', 'barS', 'bar100', 'col100')
    for i, (name, vals, color) in enumerate(series):
        ser = pl.series[i]
        if kind == 'line':
            ser.format.line.color.rgb = rgb(color)
            ser.format.line.width = Pt(line_width)
            ser.smooth = False
            from pptx.enum.chart import XL_MARKER_STYLE
            ser.marker.style = XL_MARKER_STYLE.NONE
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(color)
            ser.format.line.color.rgb = rgb('FFFFFF')
            ser.format.line.width = Pt(0.75)
            ser.invert_if_negative = False
        if point_colors and i in point_colors:
            for j, c in point_colors[i].items():
                pt = ser.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(c)
        if labels and (not isinstance(labels, (list, tuple)) or i in labels):
            dl = ser.data_labels
            dl.show_value = True
            dl.number_format = fmt
            dl.number_format_is_linked = False
            lc = (label_colors or {}).get(i, 'FFFFFF' if stacked else DARK)
            _font(dl, size, lc, bold=stacked)
            if label_pos:
                dl.position = label_pos
            elif stacked:
                dl.position = XL_LABEL_POSITION.CENTER
            elif kind != 'line':
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            if min_label is not None:
                _hide_labels(ser, [j for j, v in enumerate(vals) if v is None or abs(v) < min_label])
            elif min_share is not None:
                tot = [sum(abs(sv[j] or 0) for _, sv, _ in series) or 1 for j in range(len(cats))]
                _hide_labels(ser, [j for j, v in enumerate(vals) if v is None or abs(v) / tot[j] < min_share])
    ca = ch.category_axis
    _font(ca.tick_labels, cat_size or size)
    ca.format.line.color.rgb = rgb(LGREY)
    ca.has_major_gridlines = False
    ca.major_tick_mark = 2  # none
    if kind.startswith('bar'):
        ca.reverse_order = True
    va = ch.value_axis
    va.has_major_gridlines = grid
    if grid:
        va.major_gridlines.format.line.color.rgb = rgb('E3E5E8')
        va.major_gridlines.format.line.width = Pt(0.75)
    va.visible = value_axis
    if value_axis:
        _font(va.tick_labels, size, MUTED)
        va.format.line.fill.background()
        if axis_fmt:
            va.tick_labels.number_format = axis_fmt
            va.tick_labels.number_format_is_linked = False
    if vmax is not None:
        va.maximum_scale = vmax
    if vmin is not None:
        va.minimum_scale = vmin
    if kind.startswith('bar') and not value_axis:
        va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    if tick_skip:
        _tick_skip(ch, tick_skip)
    return ch


def legend_row(s, x, y, items, size=11, gap=0.25):
    """Легенда из цветных квадратов с подписями (для составных визуалов)."""
    cx = x
    for label, color in items:
        box(s, cx, y + 0.04, 0.16, 0.16, fill=color, radius=0.03)
        w = 0.1 + 0.075 * len(label) * size / 11
        tb(s, cx + 0.22, y - 0.02, w, 0.28, [label], size=size, color=MUTED, anchor='m')
        cx += 0.22 + w + gap
    return cx


# ───────────────────────── таблицы
NO_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'


def table(s, x, y, w, head, rows, widths, aligns=None, size=10, rh=0.3, hh=0.36, fills=None, bold_cols=(),
          head_fill=DARK, head_color=WHITE, colors=None, head_size=None):
    """Нативная таблица. fills(r, c) → цвет ячейки или None; colors(r, c) → цвет текста."""
    n, m = len(rows) + 1, len(head)
    gf = s.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(hh + rh * len(rows)))
    tbl = gf.table
    gf._element.graphic.graphicData.tbl.tblPr.find(qn('a:tableStyleId')).text = NO_STYLE
    tot = sum(widths)
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(w * cw / tot)
    aligns = aligns or ['l'] + ['r'] * (m - 1)
    for i in range(n):
        tbl.rows[i].height = Inches(hh if i == 0 else rh)
        for j in range(m):
            cell = tbl.cell(i, j)
            val = head[j] if i == 0 else rows[i - 1][j]
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if i == 0:
                bg = head_fill
            else:
                bg = (fills(i - 1, j) if fills else None) or (ZEBRA if i % 2 == 0 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(bg)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[aligns[j]]
            r = p.add_run()
            r.text = str(val)
            fc = head_color if i == 0 else ((colors(i - 1, j) if colors else None) or DARK)
            _font(r, (head_size or size) if i == 0 else size, fc, bold=(i == 0 or j in bold_cols))
    return tbl


def exec_fill(v):
    if v is None:
        return None
    return TINT_G if v >= 0.95 else TINT_A if v >= 0.75 else TINT_R


def ktg_fill(f, p):
    if f is None or p is None:
        return None
    return TINT_G if f >= p - 0.005 else TINT_A if f >= p - 0.05 else TINT_R


def cover_fill(v):
    if v is None:
        return None
    return TINT_G if v >= 0.9 else TINT_A if v >= 0.7 else TINT_R


def kpi_tile(s, x, y, w, h, value, label, sub=None, fill=WHITE, vcolor=DARK, lcolor=MUTED, vsize=26):
    box(s, x, y, w, h, fill=fill, radius=0.16)
    paras = [{'t': value, 'bold': True, 'size': vsize, 'color': vcolor, 'space': 2},
             {'t': label, 'size': 12, 'color': lcolor, 'space': 2}]
    if sub:
        paras.append({'t': sub, 'size': 10, 'color': lcolor})
    tb(s, x + 0.2, y + 0.1, w - 0.4, h - 0.2, paras, anchor='m')


def note(s, x, y, w, h, paras, fill=TINT, size=12):
    box(s, x, y, w, h, fill=fill, radius=0.14, anchor='m', paras=paras, size=size,
        margin=(0.22, 0.05, 0.22, 0.05))


def sub(s, text, y=Y0 - 0.02):
    tb(s, X0, y, W, 0.3, [text], size=12, color=MUTED)


def star_note(s):
    tb(s, X0, 6.5, 10.6, 0.3, ['* Перемещение между площадками ограничено: в обеспеченность не входит, это возможность улучшения '
                               '(запас другой площадки сверх её потребности).'], size=9, color=MUTED)


# ───────────────────────── вычисляемые выводы
T = d['total']
fact3 = {u: sum(U(u, y)['fact'] for y in FACT_YEARS) for u in UNITS}
fact3_total = sum(fact3.values())
top5 = sorted(fact3, key=lambda u: -fact3[u])[:5]
top5_share = sum(fact3[u] for u in top5) / fact3_total
pv = d['prov']['meta']
py = {r['key']: r for r in pv['byYear']}
p26, p27 = py['2026'], py['2027']
cov26 = (p26['fromStock'] + p26['fromBuy']) / p26['value']
cov27 = (p27['fromStock'] + p27['fromBuy']) / p27['value']
unc27 = uncovered(p27)
ppm = {r['key']: r for r in pv['byPpm']}
on_release27 = ppm.get('2027|onRelease', {}).get('value', 0)
never27 = ppm.get('2027|never', {}).get('value', 0)
unc27_units = sorted(((u, uncovered(prov_u(u, '2027'))) for u in UNITS if prov_u(u, '2027')), key=lambda x: -x[1])
t26 = T['2026']
t27 = T['2027']
released26 = t26['groups']['Деблокирован, пусто']
approving26 = t26['groups']['Согласование']
fleet_n = len(d['fleet'])

# ═════════════════════════ 1. Титул
s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
for ph in s.placeholders:
    if ph.placeholder_format.idx == 0:
        ph.left, ph.top, ph.width, ph.height = Inches(0.85), Inches(2.9), Inches(5.6), Inches(3.2)
        ph.text_frame.text = 'Экскаваторы WK:\nисполнение 2024–2026\nи план 2027'
        for para in ph.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(34)
    else:
        ph.left, ph.top, ph.width, ph.height = Inches(0.85), Inches(6.29), Inches(5.6), Inches(0.95)
        ph.text_frame.text = f'Данные на {AS_OF}'
s.notes_slide.notes_text_frame.text = (
    'Презентация собрана из витрин интерактивного отчёта WK CRM (ветка grok репозитория TZ): Сводка, Парк, '
    'График · план-факт, Обеспеченность, Запасы, Закупки, Кодификация, Взаимозаменяемость, Качество данных. '
    'История заказов — выгрузки BW PM-M06 2024–2027, статусы заказов — те же выгрузки, остатки — на 14.09.2026, '
    'закупки — на 11.09.2026.')

# ═════════════════════════ 2. Главное
s = content('Главное на одной странице',
            'Исполнение — факт МТР и УСО к плану заказов SAP; позиции графика ППР без заказа в план не входят, '
            'у оригиналов БЕ, перенесённых копией в «Развитие», неисполненный план не учитывается. 2026 год — на 14.09.')
tiles = [
    (f'{fleet_n}', 'бортов WK', '3 площадки · WK-20, WK-20C, WK-35', WHITE, DARK),
    (pct(T['2024']['exec']), 'исполнение 2024', f'факт {bln(T["2024"]["fact"])} из {bln(T["2024"]["plan"])} млрд ₽', WHITE, DARK),
    (pct(T['2025']['exec']), 'исполнение 2025', f'факт {bln(T["2025"]["fact"])} из {bln(T["2025"]["plan"])} млрд ₽', WHITE, DARK),
    (pct(t26['exec']), 'исполнение 2026', f'на {AS_OF}: {bln(t26["fact"])} из {bln(t26["plan"])} млрд ₽', DARK, GREEN),
    (bln(t27['plan']), 'млрд ₽ — план 2027', f'+ {bln(t27["noOrderPlan"])} млрд ₽ позиций ППР без заказа', DARK, GREEN),
    (pct(cov27), 'обеспеченность 2027', f'не покрыто {mln(unc27)} млн ₽; 2026 — {pct(cov26)}', WHITE, C_GAP),
]
tw = (W - 5 * 0.2) / 6
for i, (v, l, sb, f, vc) in enumerate(tiles):
    kpi_tile(s, X0 + i * (tw + 0.2), Y0 + 0.05, tw, 1.55, v, l, sb, fill=f, vcolor=vc,
             lcolor=WHITE if f == DARK else MUTED, vsize=28)
msgs = [
    ('2024–2025 выполнены.', f' Факт {pct(T["2024"]["exec"])} и {pct(T["2025"]["exec"])} плана; '
     f'закрыто технически и коммерчески {pct(T["2024"]["closedShare"])} и {pct(T["2025"]["closedShare"])} плана.'),
    ('2026 отстаёт от календаря.', f' К середине сентября освоено {pct(t26["exec"])}. {mln(released26)} млн ₽ '
     f'деблокированы без факта, {mln(approving26)} млн ₽ ещё на согласовании.'),
    ('Затраты сосредоточены.', f' Пять бортов — {", ".join(top5[:3])} и ещё два — дали {pct(top5_share)} факта 2024–2026.'),
    ('2027 — главный риск по МТР.', f' Не покрыто {mln(unc27)} млн ₽; {mln(on_release27)} млн ₽ закупка не видит до '
     f'деблокирования. Перемещение между площадками ограничено: запасом других площадок можно улучшить ещё до '
     f'{mln(pv["wk"]["transferPotential"])} млн ₽.'),
]
box(s, X0, Y0 + 1.85, W, 3.4, fill=WHITE, radius=0.18)
for i, (b, t) in enumerate(msgs):
    yy = Y0 + 2.0 + i * 0.8
    box(s, X0 + 0.25, yy + 0.12, 0.44, 0.44, fill=[C_FACT, C_LATE, C_BUY, C_GAP][i], radius=0.22,
        paras=[{'t': str(i + 1), 'bold': True, 'align': 'c'}], size=14, color=WHITE, anchor='m', margin=(0, 0, 0, 0))
    tb(s, X0 + 0.9, yy, W - 1.2, 0.7, [{'t': [B(b), (t, {})]}], size=14, anchor='m')

# ═════════════════════════ 3. Раздел: парк
section('01', 'Парк и готовность', f'{fleet_n} экскаваторов на трёх площадках, КТГ по годам, месяцам и бортам')

# ═════════════════════════ 4. Парк по площадкам
s = content('Парк WK по площадкам',
            'Площадка машины — по вкладке «Парк» (TOPO/ТОРО). Книга — каталог комплектации LinkOne; у четырёх бортов '
            'книга не подтверждена листом «Тех парк». КТГ — фактический коэффициент технической готовности за январь–август 2026.')
cw3 = (W - 0.6) / 3
for k, site in enumerate(SITES):
    us = [u for u in units_of(site) if u in FLEET]
    x = X0 + k * (cw3 + 0.3)
    box(s, x, Y0 + 0.05, cw3, 5.2, fill=WHITE, radius=0.18)
    box(s, x, Y0 + 0.05, cw3, 0.95, fill=DARK, radius=0.18)
    box(s, x, Y0 + 0.6, cw3, 0.4, fill=DARK)
    models = {}
    for u in us:
        models[model_of(u)] = models.get(model_of(u), 0) + 1
    tb(s, x + 0.25, Y0 + 0.1, cw3 - 1.4, 0.85, [{'t': SN[site], 'bold': True, 'size': 17, 'color': WHITE, 'space': 1},
                                                {'t': ' · '.join(f'{m} × {n}' for m, n in models.items()), 'size': 11, 'color': LGREY}],
       anchor='m')
    tb(s, x + cw3 - 1.2, Y0 + 0.1, 1.0, 0.85, [{'t': str(len(us)), 'bold': True, 'size': 30, 'color': GREEN, 'align': 'r'}],
       anchor='m')
    rh = min(0.245, 4.0 / max(1, len(us)))
    tb(s, x + 0.25, Y0 + 1.08, cw3 - 0.5, 0.25, [{'t': [('Борт · книга', {}), ]}], size=9, color=MUTED)
    tb(s, x + cw3 - 1.25, Y0 + 1.08, 1.0, 0.25, [{'t': 'КТГ 2026', 'align': 'r'}], size=9, color=MUTED)
    for j, u in enumerate(us):
        yy = Y0 + 1.33 + j * rh
        book = FLEET[u]['book'] or 'нет книги'
        tb(s, x + 0.25, yy, cw3 - 1.4, rh, [{'t': [(u, {'bold': True}), ('  ' + book, {'color': MUTED if FLEET[u]['book'] else C_GAP, 'size': 9})]}],
           size=11, anchor='m')
        kf, kp = ktg(u, '2026', 'fact'), ktg(u, '2026', 'plan')
        f = ktg_fill(kf, kp) or 'EEEFF1'
        box(s, x + cw3 - 1.05, yy + 0.025, 0.8, rh - 0.05, fill=f, radius=(rh - 0.05) / 2,
            paras=[{'t': pct(kf) if kf else 'нет данных', 'align': 'c', 'bold': bool(kf)}], size=9 if kf else 7,
            anchor='m', margin=(0, 0, 0, 0))
legend_row(s, X0, 6.5, [('КТГ факт ≥ плана', TINT_G), ('ниже плана до 5 п.п.', TINT_A), ('ниже более чем на 5 п.п.', TINT_R)], size=10)

# ═════════════════════════ 5. КТГ по месяцам
s = content('КТГ парка: план и факт по месяцам',
            'Среднее по бортам с данными, вкладка «Парк». Факт есть по август 2026 включительно. Таблица справа — '
            'среднегодовые значения по площадкам; 2026 — январь–август.')
months = d['ktg']['months']
idx = [i for i, m_ in enumerate(months) if '2024-01' <= m_ <= '2026-12']
MN = ['янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек']
cats = [f'{MN[int(months[i][5:]) - 1]} {months[i][2:4]}' for i in idx]
fm = d['ktg']['fleetMonth']
cw_ = 7.9
box(s, X0, Y0 + 0.05, cw_, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, cw_ - 0.5, 0.35, [{'t': 'Парк WK, среднее по бортам', 'bold': True}], size=14)
ch = chart(s, X0 + 0.1, Y0 + 0.5, cw_ - 0.2, 4.7, 'line', cats,
           [('КТГ план', [fm['plan'][i] for i in idx], '9AA3AB'), ('КТГ факт', [fm['fact'][i] for i in idx], C_FACT)],
           fmt='0%', labels=False, value_axis=True, grid=True, axis_fmt='0%', vmin=0.6, vmax=1.0, tick_skip=3, size=10)
rx = X0 + cw_ + 0.3
rw = W - cw_ - 0.3
box(s, rx, Y0 + 0.05, rw, 5.2, fill=WHITE, radius=0.18)
tb(s, rx + 0.25, Y0 + 0.15, rw - 0.5, 0.35, [{'t': 'По площадкам, план / факт', 'bold': True}], size=14)
rows, fl = [], []
for site in SITES:
    ks = d['ktg']['siteYear'][site]
    rows.append([SN[site]] + [f'{pct(ks[y]["plan"])} / {pct(ks[y]["fact"])}' for y in FACT_YEARS])
    fl.append([None] + [ktg_fill(ks[y]['fact'], ks[y]['plan']) for y in FACT_YEARS])
table(s, rx + 0.2, Y0 + 0.6, rw - 0.4, ['Площадка', '2024', '2025', '2026'], rows, [1.9, 1.2, 1.2, 1.2],
      aligns=['l', 'c', 'c', 'c'], size=10, rh=0.5, fills=lambda r, c: fl[r][c])
best = max(UNITS, key=lambda u: ktg(u, '2026', 'fact') or 0)
worst = min((u for u in UNITS if ktg(u, '2026', 'fact')), key=lambda u: ktg(u, '2026', 'fact'))
note(s, rx + 0.2, Y0 + 2.9, rw - 0.4, 2.2, [
    {'t': 'Разброс по бортам 2026', 'bold': True, 'space': 6},
    {'t': [('Лучший: ', {'color': MUTED}), (f'{best} — {pct(ktg(best, "2026", "fact"))}', {'bold': True})], 'space': 4},
    {'t': [('Худший: ', {'color': MUTED}), (f'{worst} — {pct(ktg(worst, "2026", "fact"))}', {'bold': True})], 'space': 6},
    {'t': 'Сухой Лог: КТГ факт есть не по всем бортам — WK-20 №1220 без данных с 2025 года.', 'size': 11, 'color': MUTED},
], size=12)

# ═════════════════════════ 6. КТГ по бортам
s = content('КТГ по каждому борту: план и факт 2024–2026',
            'Среднее по месяцам года, вкладка «Парк». Цвет ячейки факта: зелёный — не ниже плана, янтарный — ниже '
            'до 5 п.п., красный — ниже более чем на 5 п.п. «—» — нет данных в витрине КТГ.')
rows, fl = [], []
for u in [u for site in SITES for u in units_of(site) if u in FLEET]:
    r = [SN[site_of(u)].split(' ')[0], u]
    f = [None, None]
    for y in FACT_YEARS:
        p_, f_ = ktg(u, y, 'plan'), ktg(u, y, 'fact')
        r += [pct(p_), pct(f_)]
        f += [None, ktg_fill(f_, p_)]
    p_, f_ = ktg(u, '2026', 'plan'), ktg(u, '2026', 'fact')
    r.append('—' if p_ is None or f_ is None else f'{(f_ - p_) * 100:+.0f}'.replace('-', '−'))
    f.append(None)
    rows.append(r)
    fl.append(f)
table(s, X0, Y0 - 0.05, W, ['Площадка', 'Борт', '2024 план', '2024 факт', '2025 план', '2025 факт', '2026 план',
                            '2026 факт', 'Δ 2026, п.п.'],
      rows, [1.3, 1.3, 1, 1, 1, 1, 1, 1, 1.1], aligns=['l', 'l'] + ['c'] * 7, size=8.5, rh=0.18, hh=0.26,
      fills=lambda r, c: fl[r][c], bold_cols=(1,), head_size=9)

# ═════════════════════════ 7. Раздел: исполнение
section('02', 'Исполнение 2024–2026', 'План и факт по годам, площадкам, моделям, видам работ и каждому борту')

# ═════════════════════════ 8. Программа по годам
s = content('Программа ремонтов WK: план и факт по годам',
            'План и факт — МТР Полюса и МТР подрядчика (УСО) по заказам ТОРО, вкладка «График · план-факт». Позиции '
            f'графика ППР без заказа SAP в план 2027 не входят ({mln(t27["noOrderPlan"])} млн ₽, {t27["noOrderN"]} позиций). '
            f'У оригиналов БЕ с копией в «Развитии» не учтён неисполненный план: '
            f'{", ".join(mln(T[y]["copyExcluded"], 1) for y in YEARS)} млн ₽ по годам.')
cats = ['2024', '2025', f'2026 · на {AS_OF[:5]}', '2027 · план']
cw_ = 8.4
box(s, X0, Y0 + 0.05, cw_, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, cw_ - 0.5, 0.35, [{'t': 'млн ₽', 'color': MUTED}], size=12)
chart(s, X0 + 0.1, Y0 + 0.4, cw_ - 0.2, 4.8, 'col', cats,
      [('План', [T[y]['plan'] / M for y in YEARS], C_PLAN), ('Факт', [T[y]['fact'] / M if y != '2027' else None for y in YEARS], C_FACT)],
      size=12, gap=70, overlap=-10, vmax=max(T[y]['plan'] for y in YEARS) / M * 1.12)
rx = X0 + cw_ + 0.3
rw = W - cw_ - 0.3
for i, y in enumerate(FACT_YEARS):
    t = T[y]
    kpi_tile(s, rx, Y0 + 0.05 + i * 1.3, rw, 1.18, pct(t['exec']), f'исполнение {y}{" на " + AS_OF if y == "2026" else ""}',
             f'закрыто {pct(t["closedShare"])} плана · {sp(t["orders"])} заказов', fill=DARK if y == '2026' else WHITE,
             vcolor=GREEN if y == '2026' else DARK, lcolor=WHITE if y == '2026' else MUTED, vsize=26)
kpi_tile(s, rx, Y0 + 3.95, rw, 1.3, f'{mln(t27["plan"])} млн ₽', 'план 2027 в заказах SAP',
         f'{sp(t27["orders"])} заказов · МТР {mln(t27["mtrPlan"])} · УСО {mln(t27["usoPlan"])}', vsize=22)

# ═════════════════════════ 9. По площадкам
s = content('Исполнение по площадкам',
            'План и факт МТР + УСО, млн ₽. Шкалы графиков у площадок разные — сравнивайте проценты исполнения под графиками. '
            '2026 — на 14.09.2026.')
cw3 = (W - 0.6) / 3
for k, site in enumerate(SITES):
    x = X0 + k * (cw3 + 0.3)
    box(s, x, Y0 + 0.05, cw3, 5.2, fill=WHITE, radius=0.18)
    tb(s, x + 0.25, Y0 + 0.15, cw3 - 0.5, 0.4, [{'t': SN[site], 'bold': True}], size=15)
    ss = d['site'][site]
    chart(s, x + 0.05, Y0 + 0.5, cw3 - 0.1, 3.2, 'col', ['2024', '2025', '2026', '2027 план'],
          [('План', [ss[y]['plan'] / M for y in YEARS], C_PLAN), ('Факт', [ss[y]['fact'] / M if y != '2027' else None for y in YEARS], C_FACT)],
          size=10, gap=60, overlap=-10, vmax=max(ss[y]['plan'] for y in YEARS) / M * 1.18)
    for j, y in enumerate(FACT_YEARS):
        e = ss[y]['exec']
        xx = x + 0.25 + j * ((cw3 - 0.5) / 3)
        box(s, xx, Y0 + 3.85, (cw3 - 0.5) / 3 - 0.1, 0.9, fill=(exec_fill(e) if y != '2026' else 'E3E6F7') or ZEBRA, radius=0.12, anchor='m',
            paras=[{'t': pct(e), 'bold': True, 'size': 18, 'align': 'c', 'space': 0},
                   {'t': f'исполнение {y}' + (' · на 14.09' if y == '2026' else ''), 'size': 9, 'align': 'c', 'color': MUTED}], margin=(0.02, 0.02, 0.02, 0.02))
    tb(s, x + 0.25, Y0 + 4.82, cw3 - 0.5, 0.4,
       [{'t': f'Бортов: {len([u for u in units_of(site) if u in FLEET])} · заказов 2026: {sp(ss["2026"]["orders"])}', 'color': MUTED}], size=10)

# ═════════════════════════ 10. Статусы заказов по годам
s = content('Где находятся заказы: стадии SAP по годам',
            'Стадия заказа по системным и пользовательским статусам PM-06: закрыт (ЗАКР), тех. закрыт (ТЗКР, ПРСЗ, ВСБЕ), '
            'в работе (факт, подтверждения, ПРНТ, ФХСМ), деблокирован без факта и подтверждений, на согласовании '
            '(ПЛАН–СГГС). Статус — на дату годовой выгрузки. Доли — от плана в рублях.')
G = [g for g in d['groups'] if g != 'Нет статуса']
lw = 7.3
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, lw - 0.5, 0.35, [{'t': 'Все площадки, доля плана по стадиям', 'bold': True}], size=14)
chart(s, X0 + 0.1, Y0 + 0.5, lw - 0.2, 4.7, 'bar100', ['2024', '2025', '2026', '2027 план'],
      [(g, [T[y]['groups'][g] / M for y in YEARS], C_GROUP[g]) for g in G], fmt='# ##0', size=10, gap=45,
      min_label=60, label_colors={G.index('В работе'): DARK, G.index('ППР без заказа'): DARK}, legend_pos='b')
rx = X0 + lw + 0.3
rw = W - lw - 0.3
box(s, rx, Y0 + 0.05, rw, 5.2, fill=WHITE, radius=0.18)
tb(s, rx + 0.25, Y0 + 0.15, rw - 0.5, 0.35, [{'t': '2026 по площадкам, доли; подписи — млн ₽', 'bold': True}], size=14)
G26 = [g for g in G if g != 'ППР без заказа']
chart(s, rx + 0.1, Y0 + 0.5, rw - 0.2, 3.1, 'bar100', [SHORT[x] for x in SITES],
      [(g, [d['site'][x]['2026']['groups'][g] / M for x in SITES], C_GROUP[g]) for g in G26], size=9, gap=40,
      min_share=0.09, legend=False, label_colors={G26.index('В работе'): DARK})
note(s, rx + 0.2, Y0 + 3.7, rw - 0.4, 1.45, [
    {'t': [B(f'{mln(released26)} млн ₽'), (f' — {t26["groupN"]["Деблокирован, пусто"]} заказов 2026 деблокированы, но без факта и подтверждений.', {})], 'space': 4},
    {'t': [B(f'{mln(approving26)} млн ₽'), (' на согласовании в сентябре: пересмотрите, что реально выполнится до конца года.', {})]},
], size=11, fill=TINT)

# ═════════════════════════ 11. Виды работ и модели
s = content('На что уходят деньги: виды работ и модели',
            'Вид работ — по наибольшей доле плана в заказе. «Не присвоено» — заказы на компоненты (крупные узлы) без вида '
            'работ. Факт 2024–2026 и план 2027, млн ₽.')
W_ = d['work']
top_w = sorted(W_, key=lambda k: -(sum(W_[k][y][1] for y in FACT_YEARS) + W_[k]['2027'][0]))[:8]
lw = 8.1
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
chart(s, X0 + 0.1, Y0 + 0.15, lw - 0.2, 5.0, 'bar', [('Не присвоено (компоненты)' if k == 'Не присвоено' else k) for k in top_w],
      [(f'Факт {y}', [W_[k][y][1] / M for k in top_w], C_Y[y]) for y in FACT_YEARS] +
      [('План 2027', [W_[k]['2027'][0] / M for k in top_w], C_Y['2027'])],
      size=9, gap=35, overlap=0, min_label=40, legend_pos='t')
rx = X0 + lw + 0.3
rw = W - lw - 0.3
box(s, rx, Y0 + 0.05, rw, 5.2, fill=WHITE, radius=0.18)
tb(s, rx + 0.25, Y0 + 0.15, rw - 0.5, 0.35, [{'t': 'Исполнение по моделям', 'bold': True}], size=14)
rows, fl = [], []
for mdl in ['WK-20', 'WK-20C', 'WK-35']:
    mm = d['model'][mdl]
    rows.append([mdl] + [pct(mm[y]['exec']) for y in FACT_YEARS] + [mln(mm['2027']['plan'])])
    fl.append([None] + [exec_fill(mm[y]['exec']) if y != '2026' else None for y in FACT_YEARS] + [None])
table(s, rx + 0.2, Y0 + 0.6, rw - 0.4, ['Модель', '2024', '2025', '2026', 'План 27'], rows, [1.2, 1, 1, 1, 1.2],
      aligns=['l', 'c', 'c', 'c', 'r'], size=10, rh=0.42, fills=lambda r, c: fl[r][c], bold_cols=(0,))
R = d['reason']
tb(s, rx + 0.25, Y0 + 2.5, rw - 0.5, 0.35, [{'t': 'По статье затрат, факт 2024–2026', 'bold': True}], size=14)
rs = sorted(R, key=lambda k: -sum(R[k][y][1] for y in FACT_YEARS))[:3]
rtot = sum(sum(R[k][y][1] for y in FACT_YEARS) for k in R)
for j, k in enumerate(rs):
    v = sum(R[k][y][1] for y in FACT_YEARS)
    yy = Y0 + 2.95 + j * 0.74
    tb(s, rx + 0.25, yy, rw - 0.5, 0.3, [{'t': [(k, {'bold': True})]}], size=10)
    box(s, rx + 0.25, yy + 0.3, max(0.05, (rw - 1.9) * v / rtot), 0.24, fill=[C_FACT, C_BUY, C_LATE][j], radius=0.05)
    tb(s, rx + 0.35 + (rw - 1.9) * v / rtot, yy + 0.26, 1.6, 0.3, [f'{mln(v)} млн ₽ · {pct(v / rtot)}'], size=10, color=MUTED)

# ═════════════════════════ 12. Затраты по бортам
s = content('Затраты на каждый борт: факт 2024–2026',
            'Факт МТР + УСО по заказам ТОРО, млн ₽, по годам. Борта отсортированы по сумме за три года. '
            f'Пять крупнейших — {pct(top5_share)} всего факта: капитальные ремонты и замена ковшей, рукоятей, ходовой.')
ord_u = [u for u in sorted(UNITS, key=lambda u: -fact3[u]) if fact3[u] > 0]
box(s, X0, Y0 + 0.05, W, 5.25, fill=WHITE, radius=0.18)
chart(s, X0 + 0.1, Y0 + 0.1, W - 0.2, 5.15, 'barS', [f'{u} · {SHORT[site_of(u)]}' for u in ord_u],
      [(f'Факт {y}', [U(u, y)['fact'] / M for u in ord_u], C_Y[y]) for y in FACT_YEARS], size=8, gap=35,
      min_label=45, label_colors={0: DARK, 1: DARK}, legend_pos='t', cat_size=8.5)


# ═════════════════════════ 13–14. Таблицы по бортам
def unit_table(title, sites, notes):
    s = content(title, notes)
    rows, fl = [], []
    for site in sites:
        for u in units_of(site):
            r = [SHORT[site], u]
            f = [None, None]
            for y in FACT_YEARS:
                uu = U(u, y)
                r += [mln(uu['plan']) if uu['plan'] else '—', mln(uu['fact']) if uu['fact'] else '—', pct(uu['exec']) if uu['plan'] else '—']
                f += [None, None, (exec_fill(uu['exec']) if y != '2026' else None) if uu['plan'] else None]
            u27 = U(u, '2027')
            r += [mln(u27['plan']) if u27['plan'] else '—', mln(u27['noOrderPlan']) if u27['noOrderPlan'] else '—']
            f += [None, None]
            p = prov_u(u, '2027')
            cv = (p['fromStock'] + p['fromBuy']) / p['value'] if p and p['value'] else None
            r.append(pct(cv))
            f.append(cover_fill(cv))
            kf = ktg(u, '2026', 'fact')
            r.append(pct(kf))
            f.append(ktg_fill(kf, ktg(u, '2026', 'plan')))
            rows.append(r)
            fl.append(f)
    rh = min(0.28, 4.95 / len(rows))
    table(s, X0, Y0 - 0.05, W,
          ['Площадка', 'Борт', 'План 24', 'Факт 24', '%', 'План 25', 'Факт 25', '%', 'План 26', 'Факт 26', '%',
           'План 27', 'ППР 27*', 'Обесп. 27', 'КТГ 26'],
          rows, [1.15, 1.2, 0.8, 0.8, 0.62, 0.8, 0.8, 0.62, 0.8, 0.8, 0.62, 0.8, 0.8, 0.85, 0.75],
          aligns=['l', 'l'] + ['r', 'r', 'c'] * 3 + ['r', 'r', 'c', 'c'], size=9, rh=rh, hh=0.3,
          fills=lambda r, c: fl[r][c], bold_cols=(1,), head_size=9,
          colors=lambda r, c: MUTED if c == 12 else None)
    tb(s, X0, 6.55, W, 0.3, ['млн ₽. * ППР 27 — позиции графика ППР без заказа SAP, в план не входят. Обесп. 27 — доля потребности МТР 2027, '
                             'покрытая складом и закупкой к сроку. КТГ 26 — факт январь–август.'], size=9, color=MUTED)


unit_table('Исполнение по бортам: Красноярск / Еруда', ['1100'],
           'Каждый борт площадки 1100. Цвет % исполнения 2024–2025: зелёный ≥ 95%, янтарный 75–95%, красный < 75%. '
           '2026 — на 14.09.2026 и не окрашен: год не завершён, сравнивайте борта между собой.')
unit_table('Исполнение по бортам: Магадан и Сухой Лог', ['1400', '2400'],
           'Борта площадок 1400 и 2400. WK-35 №1222–1224 и №1228 в Сухом Логе — новые машины: заказов SAP нет, '
           'в 2027 есть только позиции графика ППР. №1228 нет в витрине «Парк».')

# ═════════════════════════ 15. Крупнейшие материалы
s = content('Крупнейшие материалы: факт 2024–2026',
            'Факт МТР по коду ЕКМТР за три года, млн ₽, вкладка «График · план-факт». Полоса — доля от крупнейшей позиции. '
            'Первые позиции — ковши, канаты, ходовая и смазки: они и определяют бюджет.')
tf_ = d['topFact'][:12]
mx = tf_[0]['fact']
rows = [[str(i + 1), x['code'], x['name'][:48], sp(x['qf']), mln(x['fact']), ''] for i, x in enumerate(tf_)]
rh = 0.38
tbl_w = W
table(s, X0, Y0 + 0.05, tbl_w, ['№', 'ЕКМТР', 'Наименование', 'Кол-во', 'Факт, млн ₽', ''], rows,
      [0.35, 0.9, 5.0, 0.9, 1.1, 3.0], aligns=['c', 'l', 'l', 'r', 'r', 'l'], size=10, rh=rh, hh=0.34, bold_cols=(4,))
bx = X0 + tbl_w * (0.35 + 0.9 + 5.0 + 0.9 + 1.1) / 11.25 + 0.1
bw = tbl_w * 3.0 / 11.25 - 0.25
for i, x in enumerate(tf_):
    box(s, bx, Y0 + 0.05 + 0.34 + i * rh + 0.1, bw * x['fact'] / mx, rh - 0.2, fill=C_FACT if i < 3 else '7FD8B4', radius=0.04)

# ═════════════════════════ 16. Раздел: план 2027
section('03', 'План 2027', 'Объём, готовность заказов, разрез по площадкам и бортам, крупнейшие МТР')

# ═════════════════════════ 17. Структура плана 2027
s = content('План 2027: объём и готовность заказов',
            'План 2027 по стадиям SAP на дату выгрузки, млн ₽. Позиции графика ППР без заказа (номера 1ХХХХХХХХХХХ) — '
            'не заказы: в план и потребность не включаются, показаны отдельно, чтобы был виден полный объём графика.')
G27 = ['Деблокирован, пусто', 'Согласование', 'ППР без заказа']
lw = 7.6
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, lw - 0.5, 0.35, [{'t': 'По площадкам, млн ₽', 'bold': True}], size=14)
chart(s, X0 + 0.1, Y0 + 0.5, lw - 0.2, 4.7, 'colS', [SN[x] for x in SITES],
      [({'Деблокирован, пусто': 'Деблокирован', 'Согласование': 'Открыт: согласование и СГГС', 'ППР без заказа': 'Позиции ППР без заказа'}[g],
        [d['site'][x]['2027']['groups'][g] / M for x in SITES], C_GROUP[g]) for g in G27],
      size=11, gap=55, min_label=40, label_colors={2: DARK})
rx = X0 + lw + 0.3
rw = W - lw - 0.3
kpi_tile(s, rx, Y0 + 0.05, rw, 1.25, f'{mln(t27["plan"])} млн ₽', 'план 2027 в заказах SAP',
         f'{sp(t27["orders"])} заказов · из них {pct(t27["groups"]["Согласование"] / t27["plan"])} ещё открыты', vsize=24)
kpi_tile(s, rx, Y0 + 1.45, rw, 1.25, f'{mln(t27["noOrderPlan"])} млн ₽', 'позиции ППР без заказа',
         f'{sp(t27["noOrderN"])} позиций · не входят в план', fill=ZEBRA, vsize=24)
kpi_tile(s, rx, Y0 + 2.85, rw, 1.25, f'{mln(t27["usoPlan"])} млн ₽', 'из плана — МТР подрядчика (УСО)',
         f'Магадан: {mln(d["site"]["1400"]["2027"]["usoPlan"])} млн ₽', vsize=24)
note(s, rx, Y0 + 4.25, rw, 1.0, [{'t': [B('Деблокировано '), (f'{mln(t27["groups"]["Деблокирован, пусто"])} млн ₽ — '
                                                              'остальное закупка увидит только после ДЕБЛ, если у строк «Начиная с деблок.».', {})]}],
     size=11)

# ═════════════════════════ 18. План 2027 по бортам
s = content('План 2027 по бортам',
            'План в заказах SAP и позиции графика ППР без заказа, млн ₽. Сортировка по плану в заказах. '
            'У новых бортов Сухого Лога заказов нет — только позиции ППР.')
ord27 = sorted([u for u in UNITS if U(u, '2027')['plan'] + U(u, '2027')['noOrderPlan'] > 0], key=lambda u: -(U(u, '2027')['plan'] * 10 + U(u, '2027')['noOrderPlan']))
box(s, X0, Y0 + 0.05, W, 5.25, fill=WHITE, radius=0.18)
chart(s, X0 + 0.1, Y0 + 0.1, W - 0.2, 5.15, 'barS', [f'{u} · {SHORT[site_of(u)]}' for u in ord27],
      [('План в заказах SAP', [U(u, '2027')['plan'] / M for u in ord27], C_Y['2027']),
       ('Позиции ППР без заказа', [U(u, '2027')['noOrderPlan'] / M for u in ord27], C_PLAN)],
      size=8, gap=35, min_label=25, label_colors={1: DARK}, legend_pos='t', cat_size=8.5)

# ═════════════════════════ 19. МТР 2027
s = content('Крупнейшие МТР плана 2027 и их обеспеченность',
            'План МТР 2027 по коду ЕКМТР (без позиций ППР), млн ₽; бортов — сколько машин несут потребность. '
            'Обеспечено — склад и закупка к сроку по расчёту вкладки «Обеспеченность» (2026–2027, количество).')
rows, fl = [], []
PI = d.get('provItems', {})
for i, x in enumerate(d['topPlan27'][:12]):
    pi = PI.get(str(x['code']))
    cv = ((pi['fromStock'] + pi['fromBuy']) / pi['needQty']) if pi and pi['needQty'] else None
    verdict = {'covered': 'обеспечено', 'inTime': 'успеем заказать', 'late': 'не успеем', 'past': 'срок прошёл',
               'nodate': 'нет срока'}.get(pi['verdict'], pi['verdict'] or '—') if pi else '—'
    rows.append([str(i + 1), x['code'], x['name'][:44], str(x['units']), sp(x['qp27']), mln(x['plan27']), pct(cv), verdict])
    fl.append([None] * 6 + [cover_fill(cv), {'не успеем': TINT_R, 'срок прошёл': TINT_R, 'успеем заказать': TINT_A,
                                            'обеспечено': TINT_G}.get(verdict)])
table(s, X0, Y0 + 0.05, W, ['№', 'ЕКМТР', 'Наименование', 'Бортов', 'Кол-во', 'План, млн ₽', 'Обеспечено', 'Вывод'],
      rows, [0.35, 0.9, 4.6, 0.7, 0.8, 1.1, 1.1, 1.6], aligns=['c', 'l', 'l', 'c', 'r', 'r', 'c', 'c'], size=10,
      rh=0.38, hh=0.34, fills=lambda r, c: fl[r][c], bold_cols=(5,))

# ═════════════════════════ 20. Раздел: обеспеченность
section('04', 'Обеспеченность МТР 2026–2027', 'Склад, закупка, дефицит — по годам, площадкам и бортам')

# ═════════════════════════ 21. Обеспеченность по годам
s = content('Обеспеченность потребности WK по годам',
            'Потребность открытых заказов по номенклатуре WK после правил SAP, вкладка «Обеспеченность», остатки на 14.09.2026. '
            'Сначала резервируется склад, затем закупка с датой поставки до потребности. «Опаздывает» — закупка придёт позже; '
            '«без срока» — закупка без даты или просрочена.')
star_note(s)
# Перемещение между площадками ограничено и в обеспеченность НЕ входит.
# Часть непокрытого к сроку, которую мог бы закрыть запас другой площадки,
# выделена светлым сегментом «*» — это возможность улучшения, а не покрытие.
C_POT = 'F2B8BD'
SEG = [('Склад своей площадки', 'fromStock', C_STOCK), ('Закупка к сроку', 'fromBuy', C_BUY),
       ('Закупка опаздывает / без срока', 'lateNet', C_LATE),
       ('Не покрыто, есть на другой площадке*', 'pot', C_POT), ('Не покрыто', 'gapNet', C_GAP)]
SEG_DARK = {2: DARK, 3: DARK}


def seg_vals(p, key):
    pot = p.get('transferPotential', 0.0)
    pot_gap = min(pot, p['gap'])
    if key == 'pot':
        return pot
    if key == 'gapNet':
        return p['gap'] - pot_gap
    if key == 'lateNet':
        return p['late'] + p['undated'] - (pot - pot_gap)
    return p.get(key, 0.0)


def cov_pot(p):
    return (p['fromStock'] + p['fromBuy'] + p.get('transferPotential', 0.0)) / p['value'] if p['value'] else None


lw = 8.2
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, lw - 0.5, 0.35, [{'t': 'Потребность, млн ₽', 'bold': True}], size=14)
chart(s, X0 + 0.1, Y0 + 0.5, lw - 0.2, 4.7, 'barS', [f'2026 · {mln(p26["value"])} млн ₽', f'2027 · {mln(p27["value"])} млн ₽'],
      [(n, [seg_vals(p26, k) / M, seg_vals(p27, k) / M], c) for n, k, c in SEG], size=12, gap=45, min_label=30,
      label_colors=SEG_DARK)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
kpi_tile(s, rx, Y0 + 0.05, rw, 1.25, pct(cov26), 'обеспечено к сроку 2026',
         f'не покрыто {mln(uncovered(p26))} млн ₽ · с перемещением* до {pct(cov_pot(p26))}', vsize=28)
kpi_tile(s, rx, Y0 + 1.45, rw, 1.25, pct(cov27), 'обеспечено к сроку 2027',
         f'не покрыто {mln(unc27)} млн ₽ · с перемещением* до {pct(cov_pot(p27))}', fill=DARK, vcolor=GREEN, lcolor=WHITE, vsize=28)
fz = pv['feasible']
note(s, rx, Y0 + 2.9, rw, 2.35, [
    {'t': 'Что ещё можно успеть', 'bold': True, 'space': 6},
    {'t': [('Успеем, если заказать сейчас: ', {}), B(f'{mln(fz["inTime"]["value"])} млн ₽')], 'bullet': True, 'bcolor': C_STOCK, 'space': 3},
    {'t': [('Опоздаем до 3 мес.: ', {}), B(f'{mln(fz["late3"]["value"])} млн ₽')], 'bullet': True, 'bcolor': C_LATE, 'space': 3},
    {'t': [('Опоздаем более чем на 3 мес.: ', {}), B(f'{mln(fz["lateMore"]["value"])} млн ₽')], 'bullet': True, 'bcolor': C_GAP, 'space': 3},
    {'t': 'Срок — медиана фактических поставок по коду.', 'size': 10, 'color': MUTED},
], size=11)

# ═════════════════════════ 22. По площадкам
s = content('Обеспеченность по площадкам',
            'Потребность открытых заказов WK по площадке машины, млн ₽. Таблица — доля, покрытая складом и закупкой к сроку.')
star_note(s)
cats, vals = [], {n: [] for n, _, _ in SEG}
for site in SITES:
    for y in ('2026', '2027'):
        p = d['prov']['site'][site][y]
        cats.append(f'{SN[site]} · {y}')
        for n, k, _ in SEG:
            vals[n].append(seg_vals(p, k) / M)
lw = 8.2
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
chart(s, X0 + 0.1, Y0 + 0.15, lw - 0.2, 5.05, 'barS', cats, [(n, vals[n], c) for n, _, c in SEG], size=10, gap=40,
      min_label=35, label_colors=SEG_DARK)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
rows, fl = [], []
for site in SITES:
    r, f = [SN[site]], [None]
    for y in ('2026', '2027'):
        p = d['prov']['site'][site][y]
        cv = (p['fromStock'] + p['fromBuy']) / p['value'] if p['value'] else None
        r += [pct(cv)]
        f += [cover_fill(cv)]
    r.append(mln(d['prov']['site'][site]['2026']['transferPotential'] + d['prov']['site'][site]['2027']['transferPotential']))
    f.append(None)
    r.append(mln(uncovered(d['prov']['site'][site]['2027'])))
    f.append(None)
    rows.append(r)
    fl.append(f)
table(s, rx, Y0 + 0.05, rw, ['Площадка', '2026', '2027', 'Перем.*', 'Не покр. 27'], rows, [1.5, 0.7, 0.7, 0.8, 1.0],
      aligns=['l', 'c', 'c', 'r', 'r'], size=10, rh=0.5, fills=lambda r, c: fl[r][c])
share_1100 = uncovered(d['prov']['site']['1100']['2027']) / unc27
_sl = [d['prov']['site']['2400'][y] for y in ('2026', '2027')]
sl_cov = sum(p['fromStock'] + p['fromBuy'] for p in _sl) / sum(p['value'] for p in _sl)
sl_pot = sum(p['transferPotential'] for p in _sl)
note(s, rx, Y0 + 2.3, rw, 2.95, [
    {'t': f'{pct(share_1100)} дефицита 2027 — Красноярск', 'bold': True, 'space': 6},
    {'t': f'Больше всего не покрыто у {", ".join(u for u, _ in unc27_units[:5])} — '
          f'{mln(sum(v for _, v in unc27_units[:5]))} млн ₽. Крупнейшая позиция — ковш K1839 ({mln(d["prov"]["deficit"][0]["gapValue"])} млн ₽).', 'space': 6},
    {'t': f'Сухой Лог почти без своего запаса: обеспечено {pct(sl_cov)}; перемещением* можно улучшить ещё до {mln(sl_pot)} млн ₽.',
     'color': MUTED, 'size': 11},
], size=12)

# ═════════════════════════ 23. По бортам 2027
s = content('Обеспеченность 2027 по бортам',
            'Потребность МТР 2027 открытых заказов каждого борта, млн ₽. Сортировка по непокрытой части. '
            'Борта без открытых заказов с потребностью WK не показаны.')
star_note(s)
bu = [u for u, _ in unc27_units if prov_u(u, '2027')['value'] > 0]
box(s, X0, Y0 + 0.05, W, 5.25, fill=WHITE, radius=0.18)
chart(s, X0 + 0.1, Y0 + 0.1, W - 0.2, 5.15, 'barS', [f'{u} · {SHORT[site_of(u)]}' for u in bu],
      [(n, [seg_vals(prov_u(u, '2027'), k) / M for u in bu], c) for n, k, c in SEG], size=8, gap=35, min_label=12,
      label_colors=SEG_DARK, legend_pos='t', cat_size=8.5)

# ═════════════════════════ 24. Видимость для закупки
s = content('Видит ли закупка потребность 2027',
            'Признак «Резерв./заявка» строк заказа: «Немедленно» — потребность сразу в ППМ; «Начиная с деблок.» — только '
            'после деблокирования заказа; «Никогда» — в заявку не попадает (справочные позиции). млн ₽.')
star_note(s)
PPM = [('Немедленно', 'immediate'), ('Начиная с деблок.', 'onRelease'), ('Никогда', 'never')]
lw = 7.4
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, lw - 0.5, 0.35, [{'t': '2027 по признаку, млн ₽', 'bold': True}], size=14)
pp = [ppm.get(f'2027|{k}', {'value': 0, 'fromStock': 0, 'fromBuy': 0, 'late': 0, 'undated': 0, 'gap': 0, 'transferPotential': 0}) for _, k in PPM]
chart(s, X0 + 0.1, Y0 + 0.5, lw - 0.2, 4.7, 'barS', [f'{n} · {mln(p["value"])}' for (n, _), p in zip(PPM, pp)],
      [(n, [seg_vals(p, k) / M for p in pp], c) for n, k, c in SEG], size=11, gap=45, min_label=25, label_colors=SEG_DARK)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
kpi_tile(s, rx, Y0 + 0.05, rw, 1.45, f'{mln(on_release27)} млн ₽', 'закупка не видит до деблокирования',
         f'из них не покрыто {mln(uncovered(pp[1]))} млн ₽', fill=DARK, vcolor=GREEN, lcolor=WHITE, vsize=26)
kpi_tile(s, rx, Y0 + 1.65, rw, 1.3, f'{mln(never27)} млн ₽', 'строки «Никогда»', 'в заявку не попадут совсем', vsize=24)
note(s, rx, Y0 + 3.15, rw, 2.1, [
    {'t': 'Что сделать', 'bold': True, 'space': 6},
    {'t': 'Деблокировать заказы 2027 с длинными позициями раньше — с запасом на срок поставки.', 'bullet': True, 'space': 4},
    {'t': 'Либо перевести такие строки на «Немедленно».', 'bullet': True, 'space': 4},
    {'t': 'Строки «Никогда» проверить: реальная ли это потребность.', 'bullet': True},
], size=11)

# ═════════════════════════ 25. Критичные позиции
s = content('Критичные позиции дефицита',
            'Крупнейшие по непокрытой стоимости позиции 2026–2027, вкладка «Обеспеченность». «Заказать до» — дата первой '
            'непокрытой потребности минус медианный фактический срок поставки кода.')
VERD = {'covered': ('обеспечено', TINT_G), 'inTime': ('успеем заказать', TINT_A), 'late': ('не успеем', TINT_R),
        'past': ('срок прошёл', TINT_R), 'nodate': ('нет срока', None)}


def dmy(v):
    return f'{v[8:10]}.{v[5:7]}.{v[:4]}' if v else '—'


rows, fl = [], []
for i, x in enumerate(d['prov']['deficit'][:12]):
    vl, vf = VERD.get(x['verdict'], (x['verdict'] or '—', None))
    rows.append([x['code'], x['name'][:44], sp(x['needQty']), mln(x['gapValue'], 1), str(x['leadDays'] or '—'),
                 dmy(x['firstNeed']), dmy(x['orderBy']), vl])
    fl.append([None] * 7 + [vf])
table(s, X0, Y0 + 0.05, W, ['ЕКМТР', 'Наименование', 'Потребн.', 'Не покрыто, млн ₽', 'Срок, дн.', 'Нужно с',
                            'Заказать до', 'Вывод'], rows, [0.9, 4.3, 0.85, 1.35, 0.85, 1.05, 1.1, 1.45],
      aligns=['l', 'l', 'r', 'r', 'r', 'c', 'c', 'c'], size=10, rh=0.38, hh=0.34, fills=lambda r, c: fl[r][c], bold_cols=(3,))

# ═════════════════════════ 26. Заказать сегодня
s = content('Заказать сегодня',
            'Позиции дефицита, которые ещё закрываются заказом на дату снимка: срок поставки укладывается до потребности. '
            'Это пакет «Заказать сегодня» со вкладки «Сводка» — его можно положить в заявку целиком.')
kpi_tile(s, X0, Y0 + 0.05, 3.0, 1.5, str(d['prov']['orderTodayN']), 'позиций ещё успеваем', 'если заказать сейчас', fill=DARK,
         vcolor=GREEN, lcolor=WHITE, vsize=32)
kpi_tile(s, X0, Y0 + 1.75, 3.0, 1.5, f'{mln(d["prov"]["orderTodaySum"])}', 'млн ₽ — стоимость пакета', vsize=32)
kpi_tile(s, X0, Y0 + 3.45, 3.0, 1.8, f'{mln(fz["lateMore"]["value"] + fz["late3"]["value"])}', 'млн ₽ — уже не успеть',
         'нужны аналоги, перенос сроков или ремонт узла', fill=ZEBRA, vcolor=C_GAP, vsize=32)
rows = [[x['code'], x['name'][:40], mln(x['gapValue'], 1), mln(x['canOrder'], 1), str(x['leadDays'] or '—')]
        for x in d['prov']['orderToday'][:10]]
table(s, X0 + 3.3, Y0 + 0.05, W - 3.3, ['ЕКМТР', 'Наименование', 'Не покрыто, млн ₽', 'Ещё успеем, млн ₽', 'Срок, дн.'], rows,
      [0.9, 3.8, 1.3, 1.3, 0.8], aligns=['l', 'l', 'r', 'r', 'r'], size=10, rh=0.44, hh=0.36, bold_cols=(3,))
tb(s, X0 + 3.3, Y0 + 4.95, W - 3.3, 0.3, ['«Ещё успеем» — часть дефицита, потребность по которой наступает позже срока поставки; '
                                          'остальное уже не закрыть заказом.'], size=10, color=MUTED)

# ═════════════════════════ 27. Раздел: запасы и закупки
section('05', 'Запасы, закупки и подрядчики', 'Склад WK, график поставок, сроки и МТР подрядчика (УСО)')

# ═════════════════════════ 28. Запасы по площадкам
st = d['stock']
sm = st['meta']
s = content('Запасы WK по площадкам',
            f'Остатки номенклатуры WK на {AS_OF}, вкладка «Запасы». Площадка склада — по заводу строки выгрузки MM-M03: '
            '11xx/7101/7106 — Красноярск, 14xx/7104 — Магадан, 12xx/24xx/7102/7108 — Сухой Лог (WK в Иркутской '
            'области работают только там, их заказы планирует завод 1200). Ограниченный запас вычтен на своём складе по его цене.')
star_note(s)
ss_ = st['sites']
tw4 = (W - 3 * 0.2) / 4
for i, site in enumerate(SITES):
    t = ss_.get(site, {})
    kpi_tile(s, X0 + i * (tw4 + 0.2), Y0 + 0.05, tw4, 1.3, bln(t.get('value', 0)), f'млрд ₽ — {SN[site]}',
             f'{sp(t.get("codes", 0))} кодов · {pct(t.get("value", 0) / sm["totalValue"])} запаса', vsize=26)
kpi_tile(s, X0 + 3 * (tw4 + 0.2), Y0 + 0.05, tw4, 1.3, bln(sm['totalValue']), 'млрд ₽ — весь запас WK',
         f'ограничено {mln(sm["totalRestrictedValue"], 1)} млн ₽', fill=DARK, vcolor=GREEN, lcolor=WHITE, vsize=26)
lw = 6.6
box(s, X0, Y0 + 1.6, lw, 3.65, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 1.7, lw - 0.5, 0.35, [{'t': 'Крупнейшие склады, млн ₽', 'bold': True}], size=14)
whs = st['byWarehouse'][:8]
KIND = {'transit': ' (ПБ)', 'consign': ' (конс.)'}
chart(s, X0 + 0.1, Y0 + 2.05, lw - 0.2, 3.15, 'bar', [f'{SHORT.get(site, "?")} · {n} · {pl}{KIND.get(k, "")}' for n, v, site, pl, k in whs],
      [('Запас', [v / M for n, v, site, pl, k in whs], C_BUY)], size=9, gap=35, legend=False)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
tb(s, rx, Y0 + 1.6, rw, 0.35, [{'t': 'Запас и потребность 2026–2027, млн ₽', 'bold': True}], size=14)
rows, fl = [], []
for site in SITES:
    p26, p27_ = d['prov']['site'][site]['2026'], d['prov']['site'][site]['2027']
    need = p26['value'] + p27_['value']
    own = p26['fromStock'] + p27_['fromStock']
    tr = p26['transferPotential'] + p27_['transferPotential']
    unc = uncovered(p26) + uncovered(p27_)
    rows.append([SHORT[site], mln(ss_.get(site, {}).get('value', 0)), mln(need), mln(own), mln(tr), mln(unc)])
    fl.append([None] * 4 + [TINT_A if tr > 0 else None, TINT_R if unc > 0 else None])
table(s, rx, Y0 + 2.0, rw, ['Площадка', 'Запас', 'Нужно', 'Свой склад', 'Перем.*', 'Не покр.'], rows,
      [1.5, 0.9, 0.9, 1.0, 0.95, 0.9], aligns=['l', 'r', 'r', 'r', 'r', 'r'], size=10, rh=0.5, hh=0.34,
      fills=lambda r, c: fl[r][c])
need_all = sum(d['prov']['site'][x][y]['value'] for x in SITES for y in ('2026', '2027'))
need_site = {x: sum(d['prov']['site'][x][y]['value'] for y in ('2026', '2027')) for x in SITES}
tr_site = {x: sum(d['prov']['site'][x][y]['transferPotential'] for y in ('2026', '2027')) for x in SITES}
own_sl = sum(d['prov']['site']['2400'][y]['fromStock'] for y in ('2026', '2027'))
note(s, rx, Y0 + 3.95, rw, 1.3, [
    {'t': [B('Запас лежит не там, где работы. '),
           (f'Магадан держит {pct(ss_["1400"]["value"] / sm["totalValue"])} запаса при {pct(need_site["1400"] / need_all)} потребности; '
            f'Сухой Лог со своего склада закрывает {mln(own_sl)} млн ₽ из {mln(need_site["2400"])}.', {})]}], size=11)
tw = tw4

# ═════════════════════════ 28b. Перемещения
s = content('Возможность перемещения между площадками',
            'Перемещение запаса между площадками ограничено (логистика, согласование БЕ), поэтому в обеспеченность не входит. '
            'Здесь — верхняя оценка улучшения: сколько непокрытой к сроку потребности 2026–2027 мог бы закрыть запас другой '
            'площадки, оставшийся после её собственной потребности, млн ₽. Что из этого реально переместить, решает снабжение.')
star_note(s)
moves = {(f, t_): v for f, t_, v in d['prov']['moves']}
tot_moves = sum(moves.values())
lw = 5.6
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, lw - 0.5, 0.35, [{'t': 'Откуда → куда можно, млн ₽', 'bold': True}], size=14)
rows, fl = [], []
for f in SITES:
    r = [f'из {SHORT[f]}']
    fr = [None]
    for t_ in SITES:
        v = moves.get((f, t_), 0)
        r.append('—' if f == t_ else (mln(v, 1) if v else '0'))
        fr.append(ZEBRA if f == t_ else (TINT_A if v >= 30e6 else None))
    r.append(mln(sum(moves.get((f, t_), 0) for t_ in SITES), 1))
    fr.append(None)
    rows.append(r)
    fl.append(fr)
rows.append(['всего в'] + [mln(sum(moves.get((f, t_), 0) for f in SITES), 1) for t_ in SITES] + [mln(tot_moves, 1)])
fl.append([None] * 5)
table(s, X0 + 0.2, Y0 + 0.6, lw - 0.4, ['', 'в Красноярск', 'в Магадан', 'в Сухой Лог', 'всего из'], rows,
      [1.5, 1.1, 1.0, 1.1, 0.9], aligns=['l', 'r', 'r', 'r', 'r'], size=10.5, rh=0.55, hh=0.4,
      fills=lambda r, c: fl[r][c], bold_cols=(0,))
kpi_tile(s, X0 + 0.2, Y0 + 3.55, lw - 0.4, 1.5, f'до {mln(tot_moves)} млн ₽', 'можно улучшить перемещением*',
         f'в обеспеченность не входит · 2027: {pct(cov27)} → до {pct(cov_pot(p27))}', fill=DARK, vcolor=GREEN, lcolor=WHITE, vsize=26)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
tb(s, rx, Y0 + 0.05, rw, 0.35, [{'t': 'Крупнейшие позиции, которые есть на другой площадке', 'bold': True}], size=14)
RT = {'1100': 'Красн.', '1400': 'Маг.', '2400': 'СЛ'}
rows = []
for x in d['prov']['moveTop'][:10]:
    route = ', '.join(f'{RT[a.split(">")[0]]} → {RT[a.split(">")[1]]}: {sp(q)}' for a, q in sorted(x['routes'].items(), key=lambda kv: -kv[1]))
    rows.append([x['code'], x['name'][:34], route, mln(x['value'], 1)])
table(s, rx, Y0 + 0.45, rw, ['ЕКМТР', 'Наименование', 'Маршрут, шт', 'млн ₽'], rows, [0.85, 2.9, 2.3, 0.75],
      aligns=['l', 'l', 'l', 'r'], size=9.5, rh=0.46, hh=0.34, bold_cols=(3,))

# ═════════════════════════ 29. Закупки
pu = d['purchase']
s = content('Закупки: график поставок и сроки',
            'Открытые заявки и заказы на поставку по номенклатуре WK, вкладка «Закупки», выгрузка 11.09.2026. '
            'Количества разных МТР суммированы справочно. Стоимость в выгрузке — в валюте документа (в основном ZCNY), '
            'поэтому график — в единицах. Просрочка — плановый месяц поставки раньше сентября 2026.')
tiles = [(sp(pu['items']), 'кодов в закупке', f'{sp(pu["needCodes"])} с потребностью 26–27', WHITE, DARK),
         (sp(pu['openQty']), 'ед. ещё поставить', f'в пути {sp(pu["transitQty"])} ед.', WHITE, DARK),
         (sp(pu['overdue']), 'ед. просрочено', 'месяц поставки прошёл', WHITE, C_GAP),
         (f'{pu["leadMedian"]} дн.', 'медианный срок поставки', f'по {sp(pu["leadN"])} кодам', DARK, GREEN)]
for i, (v, l, sb, f, vc) in enumerate(tiles):
    kpi_tile(s, X0 + i * (tw + 0.2), Y0 + 0.05, tw, 1.2, v, l, sb, fill=f, vcolor=vc, lcolor=WHITE if f == DARK else MUTED, vsize=24)
bm = {k: v for k, v in pu['byMonth'].items() if '2025-09' <= k <= '2027-08'}
mths = []
y_, m_ = 2025, 9
while (y_, m_) <= (2027, 8):
    mths.append(f'{y_}-{m_:02d}')
    m_ += 1
    if m_ > 12:
        y_, m_ = y_ + 1, 1
lw = 8.4
box(s, X0, Y0 + 1.45, lw, 3.8, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 1.55, lw - 0.5, 0.35, [{'t': 'Открытое количество по месяцу поставки, ед.', 'bold': True}], size=13)
chart(s, X0 + 0.1, Y0 + 1.9, lw - 0.2, 3.3, 'col', [f'{MN[int(x[5:]) - 1]} {x[2:4]}' for x in mths],
      [('Поставка', [bm.get(x, 0) for x in mths], C_BUY)], size=8, gap=40, legend=False, min_label=1,
      point_colors={0: {i: C_GAP for i, x in enumerate(mths) if x < '2026-09'}})
rx = X0 + lw + 0.3
rw = W - lw - 0.3
box(s, rx, Y0 + 1.45, rw, 3.8, fill=WHITE, radius=0.18)
tb(s, rx + 0.25, Y0 + 1.55, rw - 0.5, 0.35, [{'t': 'Фактический срок поставки, кодов', 'bold': True}], size=13)
LH = ['до 90 дн.', '90–180', '180–270', '270–360', '360–450', '450+']
chart(s, rx + 0.1, Y0 + 1.9, rw - 0.2, 3.3, 'col', LH, [('Кодов', [pu['leadHist'].get(str(i), 0) for i in range(6)], C_PLAN)],
      size=9, gap=40, legend=False, point_colors={0: {2: C_BUY, 3: C_BUY}})

# ═════════════════════════ 30. УСО
s = content('МТР подрядчика (УСО)',
            'План и факт МТР подрядчика по заказам ТОРО WK из PM-06, млн ₽; реестр — вкладка «Обеспеченность», блок УСО '
            f'({d["uso"]["meta"]["orders"]} заказов, {d["uso"]["meta"]["closedOrders"]} закрыто). 2026 — на 14.09.2026.')
lw = 7.6
box(s, X0, Y0 + 0.05, lw, 5.2, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 0.15, lw - 0.5, 0.35, [{'t': 'УСО по годам, млн ₽', 'bold': True}], size=14)
chart(s, X0 + 0.1, Y0 + 0.5, lw - 0.2, 4.7, 'col', ['2024', '2025', '2026', '2027 план'],
      [('План', [T[y]['usoPlan'] / M for y in YEARS], C_PLAN), ('Факт', [T[y]['usoFact'] / M if y != '2027' else None for y in YEARS], C_FACT)],
      size=12, gap=70, overlap=-10)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
rows = []
for site in SITES:
    ss = d['site'][site]
    rows.append([SN[site]] + [f'{mln(ss[y]["usoPlan"])} / {mln(ss[y]["usoFact"])}' for y in FACT_YEARS] + [mln(ss['2027']['usoPlan'])])
table(s, rx, Y0 + 0.05, rw, ['Площадка', '2024', '2025', '2026', '2027'], rows, [1.6, 0.9, 0.9, 0.9, 0.7],
      aligns=['l', 'c', 'c', 'c', 'r'], size=9.5, rh=0.5, hh=0.34)
tb(s, rx, Y0 + 1.95, rw, 0.3, ['план / факт, млн ₽'], size=9, color=MUTED)
note(s, rx, Y0 + 2.4, rw, 2.85, [
    {'t': 'План МТР подрядчика заполнен не полностью', 'bold': True, 'space': 6},
    {'t': f'В 2024–2025 факт УСО в {T["2024"]["usoFact"] / max(1, T["2024"]["usoPlan"]):.0f} и '
          f'{T["2025"]["usoFact"] / max(1, T["2025"]["usoPlan"]):.0f} раз выше плана строк МТР подрядчика. '
          'Процент исполнения по УСО не показателен — смотрите факт и открытый остаток.', 'space': 6},
    {'t': f'2027: {mln(d["site"]["1400"]["2027"]["usoPlan"])} млн ₽ из {mln(t27["usoPlan"])} — Магадан.', 'bold': True},
], size=11)

# ═════════════════════════ 31. Раздел: справочники
section('06', 'Справочники и качество данных', 'Каталог, кодификация, взаимозаменяемость и ограничения расчёта')

# ═════════════════════════ 32. Каталог
c_ = d['catalog']
ic = d['interchange']
s = content('Каталог, кодификация и взаимозаменяемость',
            'Вкладки «Каталог», «Каталог LinkOne», «Кодификация», «Взаимозаменяемость». Позиция без кода ЕКМТР не видна в '
            'истории расхода, остатках и закупке SAP — её нельзя обеспечить автоматически.')
cod = c_['matchedEkmtr'] / c_['items']
tiles = [(sp(c_['items']), 'позиций прайса ДП', f'УСО — {sp(c_["matchedUso"])}', WHITE, DARK),
         (pct(cod), 'кодифицировано ЕКМТР', f'{sp(c_["matchedEkmtr"])} позиций; {c_["ambiguousEkmtr"]} неоднозначно', WHITE, C_LATE),
         (sp(c_['items'] - c_['matchedEkmtr']), 'позиций без ЕКМТР', 'нет расхода, остатка и закупки', DARK, GREEN),
         (sp(c_['matchedTree']), 'позиций в дереве узлов', 'механизм → узел → деталь', WHITE, DARK)]
for i, (v, l, sb, f, vc) in enumerate(tiles):
    kpi_tile(s, X0 + i * (tw + 0.2), Y0 + 0.05, tw, 1.35, v, l, sb, fill=f, vcolor=vc, lcolor=WHITE if f == DARK else MUTED, vsize=26)
lw = 6.0
box(s, X0, Y0 + 1.65, lw, 3.6, fill=WHITE, radius=0.18)
tb(s, X0 + 0.25, Y0 + 1.75, lw - 0.5, 0.35, [{'t': 'Кодификация прайса', 'bold': True}], size=14)
chart(s, X0 + 0.1, Y0 + 2.1, lw - 0.2, 3.1, 'bar100', ['Прайс ДП'],
      [('С кодом ЕКМТР', [c_['matchedEkmtr'] - c_['ambiguousEkmtr']], C_STOCK), ('Неоднозначно', [c_['ambiguousEkmtr']], C_LATE),
       ('Без кода', [c_['items'] - c_['matchedEkmtr']], C_GAP)], size=11, gap=60, label_colors={1: DARK})
rx = X0 + lw + 0.3
rw = W - lw - 0.3
card(s, rx, Y0 + 1.65, rw, 3.6, 'Ведомость взаимозаменяемости', [
    {'t': [B(f'{sp(ic["groups"])} групп'), (f' замен, {sp(ic["partsInGroups"])} деталей', {})], 'bullet': True, 'space': 6},
    {'t': [B(f'{sp(ic["rowRelations"])} связей'), (' между книгами WK-20, WK-20C и WK-35', {})], 'bullet': True, 'space': 6},
    {'t': [B(f'{sp(ic["commentedRows"])} строк'), (' с комментарием — ограничение применения', {})], 'bullet': True, 'space': 6},
    {'t': 'Прямые замены видны в обеспеченности. Транзитивные замены и пересчёт склада «как будто замена» '
          'сознательно не делаются.', 'color': MUTED, 'size': 11},
], hfill=DARK, bsize=13)

# ═════════════════════════ 33. Ограничения
s = content('Ограничения данных и расчёта',
            'Сводка по вкладке «Качество данных» и правилам расчёта. Эти оговорки влияют на цифры презентации.')
lims = [
    ('Статус — на дату выгрузки', 'Стадия заказов 2024–2025 взята из годовых выгрузок PM-06. Если заказ закрыли позже, '
                                  'в презентации он может быть «в работе».'),
    ('Запас по площадкам', 'Площадка склада — по заводу строки остатков; склады Иркутской области отнесены к Сухому Логу. '
                           'Склад покрывает только свою площадку; перемещение ограничено и показано как возможность*, не как покрытие.'),
    ('Позиции ППР и копии', f'{mln(t27["noOrderPlan"])} млн ₽ позиций ППР 2027 и неисполненный план оригиналов БЕ '
                            'исключены из плана по правилам PM-06.'),
    ('Срок поставки — до месяца', 'Обеспеченность считается по месяцу: приход внутри месяца потребности считается успевающим.'),
    ('Парк и книги', 'В парке 28 бортов после схлопывания дублей. У 4 бортов нет книги; №1228 есть в заказах, но нет в витрине «Парк».'),
    ('Стоимость закупок', 'В выгрузке закупки стоимость в валюте документа (в основном ZCNY) — сравнивать её с рублями нельзя.'),
]
cw2, ch2 = (W - 0.6) / 3, 2.5
for i, (t, body) in enumerate(lims):
    x = X0 + (i % 3) * (cw2 + 0.3)
    y = Y0 + 0.05 + (i // 3) * (ch2 + 0.25)
    card(s, x, y, cw2, ch2, t, [body], hfill=[DARK, C_LATE, GREY, DARK, GREY, C_LATE][i],
         hcolor=DARK if [DARK, C_LATE, GREY, DARK, GREY, C_LATE][i] == C_LATE else WHITE, bsize=12, hsize=14)
tb(s, X0, 6.5, W, 0.3, [f'Отдельно: 2026 — незавершённый год (факт на {AS_OF}); низкий % исполнения 2026 не означает отставания сам по себе.'],
   size=10, color=MUTED)

# ═════════════════════════ 34. Выводы
s = content('Выводы и что сделать',
            'Предложения по итогам анализа; сроки и ответственных нужно согласовать.')
acts = [
    ('Закрыть хвосты 2026', f'{t26["groupN"]["Деблокирован, пусто"]} деблокированных заказов на {mln(released26)} млн ₽ без факта: '
                            'провести факт, перенести или закрыть до конца года.', C_LATE),
    ('Ревизия согласования 2026', f'{mln(approving26)} млн ₽ ещё на согласовании в сентябре — решить, что переходит в 2027.', C_BUY),
    ('Деблокировать 2027 раньше', f'{mln(on_release27)} млн ₽ потребности закупка не видит до ДЕБЛ при сроке поставки '
                                  f'{pv["leadMedianDays"]} дней.', C_GAP),
    ('Заказать сегодня', f'{d["prov"]["orderTodayN"]} позиций на {mln(d["prov"]["orderTodaySum"])} млн ₽ ещё успеваем; по '
                         f'{mln(fz["lateMore"]["value"] + fz["late3"]["value"])} млн ₽ — аналоги или перенос.', C_STOCK),
    ('Просроченные поставки', f'{sp(pu["overdue"])} ед. с прошедшим месяцем поставки — эскалация поставщикам.', C_GAP),
    ('Возможность перемещения', f'До {mln(pv["wk"]["transferPotential"])} млн ₽ дефицита есть на других площадках. '
                                'Перемещение ограничено — проверить, что реально согласовать и перевезти.', GREY),
]
for i, (t, body, c) in enumerate(acts):
    x = X0 + (i % 2) * ((W - 0.3) / 2 + 0.3)
    y = Y0 + 0.05 + (i // 2) * 1.75
    ww = (W - 0.3) / 2
    box(s, x, y, ww, 1.55, fill=WHITE, radius=0.18)
    box(s, x + 0.25, y + 0.45, 0.62, 0.62, fill=c, radius=0.31, paras=[{'t': str(i + 1), 'bold': True, 'align': 'c'}],
        size=18, color=WHITE, anchor='m', margin=(0, 0, 0, 0))
    tb(s, x + 1.1, y + 0.1, ww - 1.3, 1.35, [{'t': t, 'bold': True, 'size': 15, 'space': 4}, {'t': body, 'color': MUTED}],
       size=12, anchor='m')

# ═════════════════════════ 35. Методика
s = content('Методика и источники',
            'Все цифры воспроизводятся скриптами build/build_wk_status_data.py и build/build_wk_status_deck.py в репозитории TZ.')
card(s, X0, Y0 + 0.05, (W - 0.3) / 2, 5.2, 'Как считали', [
    {'t': 'Исполнение = факт / план МТР + УСО по заказам ТОРО WK.', 'bullet': True, 'space': 6},
    {'t': 'Позиции графика ППР без заказа SAP в план не входят.', 'bullet': True, 'space': 6},
    {'t': 'Оригинал БЕ с копией в «Развитии»: факт учитывается, неисполненный план — нет.', 'bullet': True, 'space': 6},
    {'t': 'Стадия заказа — по статусам PM-06: ЗАКР > ТЗКР > ДЕБЛ > ОТКР.', 'bullet': True, 'space': 6},
    {'t': 'Площадка — по борту (вкладка «Парк»).', 'bullet': True, 'space': 6},
    {'t': 'Запас — по площадке склада (завод строки остатков).', 'bullet': True, 'space': 6},
    {'t': 'Обеспеченность: склад своей площадки, затем закупка с датой до потребности, по месяцам. Запас других '
          'площадок — только возможность улучшения*, в покрытие не входит.', 'bullet': True},
], hfill=DARK, bsize=12)
card(s, X0 + (W - 0.3) / 2 + 0.3, Y0 + 0.05, (W - 0.3) / 2, 5.2, 'Источники', [
    {'t': 'Выгрузки BW PM-M06 2024–2027, площадки 1100, 1200, 1400 (2400 — по бортам)', 'bullet': True, 'space': 6},
    {'t': 'Статусы заказов PM-06 (TOPO/pm06_meta)', 'bullet': True, 'space': 6},
    {'t': 'КТГ — витрина TOPO ktg.json', 'bullet': True, 'space': 6},
    {'t': 'Остатки и ограниченный запас на 14.09.2026', 'bullet': True, 'space': 6},
    {'t': 'Закупка ALL «Развитие» на 11.09.2026', 'bullet': True, 'space': 6},
    {'t': 'МТР подрядчика — rawdata/УСО; прайсы ДП и УСО; ведомость взаимозаменяемости', 'bullet': True},
], hfill=GREEN, hcolor=DARK, bsize=13)

# ═════════════════════════ 36. Финал
s = prs.slides.add_slide(prs.slide_layouts[L_END])
for ph in s.placeholders:
    i = ph.placeholder_format.idx
    if i == 13:
        ph.text_frame.text = 'Вопросы и обсуждение'
    elif i == 14:
        ph.text_frame.text = f'Экскаваторы WK · данные на {AS_OF}'
        _nobullet(ph.text_frame.paragraphs[0])

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
