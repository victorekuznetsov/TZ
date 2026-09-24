# -*- coding: utf-8 -*-
"""Общие помощники для презентаций на шаблоне АО «Развитие» (python-pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

prs = None


def open_template(src):
    """Открыть шаблон и убрать из него слайды-примеры, оставив мастер и макеты."""
    global prs
    prs = Presentation(src)
    sld = prs.slides._sldIdLst
    for s in list(sld):
        prs.part.drop_rel(s.rId)
        sld.remove(s)
    return prs


DARK, GREEN, GREY, LGREY = '2A3138', '3EF0AF', '80868B', 'D8D8D8'
WHITE, MUTED, TINT = 'FFFFFF', '5B6167', 'D5FBEC'
FONT = 'Arial'

# Макеты шаблона
L_TITLE, L_CONTENT, L_SECTION, L_END = 10, 12, 14, 15

# Поле контента внутренних слайдов (дюймы)
X0, X1, Y0, Y1 = 0.59, 12.74, 1.18, 6.5
W = X1 - X0



def rgb(h):
    return RGBColor.from_string(h)


def _bullet(p, color=GREEN, char='•', indent=0.2):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent))))
    pPr.set('indent', str(-int(Inches(indent))))
    bc = etree.SubElement(pPr, qn('a:buClr'))
    etree.SubElement(bc, qn('a:srgbClr')).set('val', color)
    bf = etree.SubElement(pPr, qn('a:buFont'))
    bf.set('typeface', FONT)
    etree.SubElement(pPr, qn('a:buChar')).set('char', char)


def _nobullet(p):
    pPr = p._p.get_or_add_pPr()
    etree.SubElement(pPr, qn('a:buNone'))


def fill_tf(tf, paras, size=14, color=DARK, bold=False, align='l', space=4,
            anchor='t', margin=(0.0, 0.0, 0.0, 0.0), line=None):
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for item in paras:
        if not isinstance(item, dict):
            item = {'t': item}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        a = item.get('align', align)
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[a]
        p.space_after = Pt(item.get('space', space))
        if line or item.get('line'):
            p.line_spacing = item.get('line', line)
        runs = item['t']
        if isinstance(runs, str):
            runs = [(runs, {})]
        for text, st in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(st.get('size', item.get('size', size)))
            f.bold = st.get('bold', item.get('bold', bold))
            f.italic = st.get('italic', False)
            f.color.rgb = rgb(st.get('color', item.get('color', color)))
        if item.get('bullet'):
            _bullet(p, item.get('bcolor', GREEN))
        else:
            _nobullet(p)
    return tf


def tb(s, x, y, w, h, paras, **kw):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(sh.text_frame, paras, **kw)
    return sh


def box(s, x, y, w, h, fill=None, line=None, radius=0.0, shape=None, paras=None, lw=1.25, **kw):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if paras is not None:
        kw.setdefault('margin', (0.14, 0.08, 0.14, 0.08))
        fill_tf(sh.text_frame, paras, **kw)
    return sh


def chevron(s, x, y, w, h, fill, paras, first=False, **kw):
    st = MSO_SHAPE.PENTAGON if first else MSO_SHAPE.CHEVRON
    sh = box(s, x, y, w, h, fill=fill, shape=st, paras=paras, anchor='m',
             margin=(0.12 if first else h * 0.42, 0.02, h * 0.3, 0.02), **kw)
    sh.adjustments[0] = 0.32
    return sh


def card(s, x, y, w, h, head, body, hfill=DARK, hcolor=WHITE, bfill=WHITE, hsize=15,
         bsize=13, hh=0.5, space=5):
    """Карточка в стиле шаблона: цветная шапка и белое тело."""
    box(s, x, y, w, hh + 0.35, fill=hfill, radius=0.16)
    box(s, x, y + hh, w, h - hh, fill=bfill, radius=0.16)
    tb(s, x + 0.1, y, w - 0.2, hh, [{'t': head, 'bold': True, 'align': 'c'}],
       size=hsize, color=hcolor, anchor='m')
    return tb(s, x + 0.2, y + hh + 0.14, w - 0.4, h - hh - 0.2, body, size=bsize, space=space)


def content(title, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = title
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.name = FONT
        else:
            ph._element.getparent().remove(ph._element)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def section(num, title, desc, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
    for ph in s.placeholders:
        i = ph.placeholder_format.idx
        if i == 0:
            ph.text_frame.text = num
        elif i == 13:
            ph.text_frame.text = title
        elif i == 14:
            ph.text_frame.text = desc
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def source(s, text):
    tb(s, X0, 6.62, 9.6, 0.3, [text], size=10, color=MUTED)


def B(t):
    return (t, {'bold': True})
