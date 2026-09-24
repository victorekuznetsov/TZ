# -*- coding: utf-8 -*-
"""Учебная методичка «SAP ТОРО для планировщика» на шаблоне АО «Развитие».

Запуск: python3 build/build_metodichka.py <Развитие шаблон.pptx> <выход.pptx>
Шаблон лежит в TOPO, ветка rawdata. Нужен python-pptx.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

SRC = sys.argv[1] if len(sys.argv) > 1 else 'template.pptx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'SAP_TORO_metodichka_planirovshika.pptx'

DARK, GREEN, GREY, LGREY = '2A3138', '3EF0AF', '80868B', 'D8D8D8'
WHITE, MUTED, TINT = 'FFFFFF', '5B6167', 'D5FBEC'
FONT = 'Arial'

# Макеты шаблона
L_TITLE, L_CONTENT, L_SECTION, L_END = 10, 12, 14, 15

# Поле контента внутренних слайдов (дюймы)
X0, X1, Y0, Y1 = 0.59, 12.74, 1.18, 6.5
W = X1 - X0

prs = Presentation(SRC)
# убираем примеры шаблона, оставляя мастер и макеты
sld = prs.slides._sldIdLst
for s in list(sld):
    prs.part.drop_rel(s.rId)
    sld.remove(s)


def rgb(h):
    return RGBColor.from_string(h)


def _bullet(p, color=GREEN, char='•', indent=0.2):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent))))
    pPr.set('indent', str(-int(Inches(indent))))
    bc = etree.SubElement(pPr, qn('a:buClr'))
    etree.SubElement(bc, qn('a:srgbClr')).set('val', color)
    bf = etree.SubElement(pPr, qn('a:buFont'))
    bf.set('typeface', FONT)
    etree.SubElement(pPr, qn('a:buChar')).set('char', char)


def _nobullet(p):
    pPr = p._p.get_or_add_pPr()
    etree.SubElement(pPr, qn('a:buNone'))


def fill_tf(tf, paras, size=14, color=DARK, bold=False, align='l', space=4,
            anchor='t', margin=(0.0, 0.0, 0.0, 0.0), line=None):
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    first = True
    for item in paras:
        if not isinstance(item, dict):
            item = {'t': item}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        a = item.get('align', align)
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[a]
        p.space_after = Pt(item.get('space', space))
        if line or item.get('line'):
            p.line_spacing = item.get('line', line)
        runs = item['t']
        if isinstance(runs, str):
            runs = [(runs, {})]
        for text, st in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(st.get('size', item.get('size', size)))
            f.bold = st.get('bold', item.get('bold', bold))
            f.italic = st.get('italic', False)
            f.color.rgb = rgb(st.get('color', item.get('color', color)))
        if item.get('bullet'):
            _bullet(p, item.get('bcolor', GREEN))
        else:
            _nobullet(p)
    return tf


def tb(s, x, y, w, h, paras, **kw):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(sh.text_frame, paras, **kw)
    return sh


def box(s, x, y, w, h, fill=None, line=None, radius=0.0, shape=None, paras=None, lw=1.25, **kw):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if paras is not None:
        kw.setdefault('margin', (0.14, 0.08, 0.14, 0.08))
        fill_tf(sh.text_frame, paras, **kw)
    return sh


def chevron(s, x, y, w, h, fill, paras, first=False, **kw):
    st = MSO_SHAPE.PENTAGON if first else MSO_SHAPE.CHEVRON
    sh = box(s, x, y, w, h, fill=fill, shape=st, paras=paras, anchor='m',
             margin=(0.12 if first else h * 0.42, 0.02, h * 0.3, 0.02), **kw)
    sh.adjustments[0] = 0.32
    return sh


def card(s, x, y, w, h, head, body, hfill=DARK, hcolor=WHITE, bfill=WHITE, hsize=15,
         bsize=13, hh=0.5, space=5):
    """Карточка в стиле шаблона: цветная шапка и белое тело."""
    box(s, x, y, w, hh + 0.35, fill=hfill, radius=0.16)
    box(s, x, y + hh, w, h - hh, fill=bfill, radius=0.16)
    tb(s, x + 0.1, y, w - 0.2, hh, [{'t': head, 'bold': True, 'align': 'c'}],
       size=hsize, color=hcolor, anchor='m')
    return tb(s, x + 0.2, y + hh + 0.14, w - 0.4, h - hh - 0.2, body, size=bsize, space=space)


def content(title, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[L_CONTENT])
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = title
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.name = FONT
        else:
            ph._element.getparent().remove(ph._element)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def section(num, title, desc, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
    for ph in s.placeholders:
        i = ph.placeholder_format.idx
        if i == 0:
            ph.text_frame.text = num
        elif i == 13:
            ph.text_frame.text = title
        elif i == 14:
            ph.text_frame.text = desc
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def source(s, text):
    tb(s, X0, 6.62, 9.6, 0.3, [text], size=10, color=MUTED)


def B(t):
    return (t, {'bold': True})


# ───────────────────────── 1. Титул
s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
for ph in s.placeholders:
    if ph.placeholder_format.idx == 0:
        ph.left, ph.top, ph.width, ph.height = Inches(0.85), Inches(2.88), Inches(6.2), Inches(3.23)
        ph.text_frame.text = 'SAP ТОРО\nдля планировщика'
        for para in ph.text_frame.paragraphs:
            for r in para.runs:
                r.font.size = Pt(44)
    else:
        ph.left, ph.top, ph.width, ph.height = Inches(0.85), Inches(6.29), Inches(6.2), Inches(0.95)
        ph.text_frame.text = 'Учебная методичка · 2026'
s.notes_slide.notes_text_frame.text = (
    'Методичка для планировщиков ТОиР АО «Развитие». Собрана по конструкторским документам '
    'ERP.13.9.2.x, регламентам ТОиР, памяткам ТОиР-72…96 и анализу выгрузок BW PM-M06 '
    'за 2022–2027 годы (422 774 заказа).')

# ───────────────────────── 2. Содержание
s = content('Как устроена методичка',
            'Шесть модулей идут в порядке жизни заказа: от того, на каком заводе его создать, '
            'до закрытия и анализа выгрузки. В конце — чек-лист, самопроверка и справочник транзакций.')
mods = [
    ('01', 'Основы', 'Планирующий завод, оргструктура, кто за что отвечает'),
    ('02', 'Жизненный цикл заказа', 'Фазы ОТКР → ДЕБЛ → ТЗКР → ЗАКР и согласование ПЛАН → СГГС'),
    ('03', 'Выполнение и закрытие', 'Цепочка «Развития»: ПРНТ, ФХСМ, ПРСЗ, ВСБЕ и контроль в IW38'),
    ('04', 'Материалы', 'Признак «Резерв./заявка», срочность, невостребованные остатки'),
    ('05', 'Перенос и отмена работ', 'Копии заказов в «Развитие», блокировки сроков, ЗНК'),
    ('06', 'Практика', 'Чтение PM-06, типичные ошибки, чек-лист и самопроверка'),
]
cw, ch, gx, gy = (W - 0.6) / 3, 2.2, 0.3, 0.3
for i, (n, t, d) in enumerate(mods):
    x = X0 + (i % 3) * (cw + gx)
    y = Y0 + 0.2 + (i // 3) * (ch + gy)
    box(s, x, y, cw, ch, fill=WHITE, radius=0.18)
    box(s, x + 0.25, y + 0.28, 0.9, 0.9, fill=DARK if i % 2 == 0 else GREEN, radius=0.45,
        paras=[{'t': n, 'bold': True, 'align': 'c'}], size=20,
        color=GREEN if i % 2 == 0 else DARK, anchor='m', margin=(0, 0, 0, 0))
    tb(s, x + 1.35, y + 0.28, cw - 1.55, 0.9, [{'t': t, 'bold': True}], size=17, anchor='m')
    tb(s, x + 0.25, y + 1.35, cw - 0.5, 0.8, [d], size=13, color=MUTED)
source(s, 'Источники: КД ERP.13.9.2.x, регламенты ТОиР, памятки ТОиР-72…96, выгрузки BW PM-M06 2022–2027')

# ───────────────────────── 3. Раздел 01
section('01', 'Основы', 'На каком заводе создаётся заказ и кто за него отвечает')

# ───────────────────────── 4. Главное правило
s = content('Главное правило: завод определяют ресурсы',
            'Это правило из памяток ТОиР-77 и ТОиР-79. Ответственные (планирующий завод, группа '
            'плановиков, ответственное рабочее место) дальше наследуются ППР, сообщениями и заказами, '
            'поэтому ошибка на старте тянется по всей цепочке.')
box(s, X0, Y0 + 0.1, W, 0.95, fill=DARK, radius=0.18, anchor='m',
    paras=[{'t': [('Чьими ресурсами выполняются работы, ', {}),
                  ('тот планирующий завод и ставится', {'color': GREEN, 'bold': True}),
                  (' в сообщении, ППР и заказе ТОРО', {})]}],
    size=20, color=WHITE, margin=(0.35, 0.05, 0.3, 0.05))
cy, chh = Y0 + 1.35, 2.75
card(s, X0, cy, (W - 0.3) / 2, chh, 'Работы выполняет АО «Развитие»', [
    {'t': 'Сообщения и заказы — с планирующим заводом филиала «Развития»', 'bullet': True},
    {'t': 'ППР на обходы переносятся в «Развитие»: обходы делает ремонтный персонал', 'bullet': True},
    {'t': 'Давальческие МТР списываются со специального склада завода БЕ: в компонентах заказа меняют завод и склад', 'bullet': True},
], hfill=GREEN, hcolor=DARK, bsize=14)
card(s, X0 + (W - 0.3) / 2 + 0.3, cy, (W - 0.3) / 2, chh, 'Работы выполняет производственная БЕ', [
    {'t': 'Документы остаются с планирующим заводом БЕ и в «Развитие» не переносятся', 'bullet': True, 'bcolor': GREY},
    {'t': 'ППР на ЕТО остаются в БЕ: ЕТО выполняет эксплуатация', 'bullet': True, 'bcolor': GREY},
    {'t': 'Технологические карты не переносятся: в ППР «Развития» используются карты БЕ, рабочие места подставляет настройка MEBOR', 'bullet': True, 'bcolor': GREY},
], bsize=14)
box(s, X0, cy + chh + 0.25, W, 0.72, fill=TINT, radius=0.14, anchor='m',
    paras=[{'t': [B('Недельный отзыв ППР. '),
                  ('По группе «Механика», включая ЕТО, отзывы для АО «Полюс Вернинское» и АО «Развитие» '
                   'выполняют планировщики «Развития»: ремонтный персонал механики переведён в «Развитие».', {})]}],
    size=13, margin=(0.25, 0.05, 0.25, 0.05))

# ───────────────────────── 5. Оргструктура
s = content('Оргструктура ТОРО в SAP',
            'Три уровня оргструктуры. Завод расположения — где стоит техника. Завод планирования — где '
            'планируют работы и МТР и собирают затраты. Группа плановиков разграничивает полномочия '
            'внутри завода планирования. Коды заводов: ХХNN, где ХХ — первые два символа кода БЕ.')
steps = [
    ('Завод расположения', 'Где физически стоит оборудование: ГТК, ЗИФ, управляющие офисы, перевалочные базы'),
    ('Завод планирования ТОРО', 'Где планируют работы и потребность в МТР, ведут остатки и собирают затраты. Один к одному с логистическим заводом'),
    ('Группа плановиков', 'Подразделение внутри завода: механика, энергетика, КИПиА. По ней разграничены полномочия на заказы'),
]
sw = (W - 0.2) / 3
for i, (t, d) in enumerate(steps):
    x = X0 + i * (sw + 0.1)
    chevron(s, x, Y0 + 0.2, sw, 0.95, [DARK, GREY, GREEN][i],
            [{'t': t, 'bold': True, 'align': 'c'}], first=(i == 0), size=16,
            color=DARK if i == 2 else WHITE)
    tb(s, x + 0.15, Y0 + 1.35, sw - 0.5, 1.6, [d], size=14)
box(s, X0, 4.2, W, 2.15, fill=WHITE, radius=0.18)
tb(s, X0 + 0.3, 4.35, 3.2, 0.4, [{'t': 'Откуда заказ берёт ответственных', 'bold': True}], size=15)
flow = ['ТМ / ЕО', 'План ППР', 'Сообщение', 'Заказ ТОРО']
fw = 2.35
for i, t in enumerate(flow):
    chevron(s, X0 + 0.3 + i * (fw + 0.05), 4.85, fw, 0.62, LGREY if i < 3 else DARK,
            [{'t': t, 'bold': True, 'align': 'c'}], first=(i == 0), size=13,
            color=DARK if i < 3 else WHITE)
tb(s, X0 + 0.3, 5.62, W - 0.6, 0.7, [
    {'t': [B('После отзыва ППР ответственных в заказе изменить нельзя. '),
           ('Массово сменить ответственных можно заранее: в ТМ — IL05, в ЕО — IE05, в позициях ППР — IP17.', {})]}],
   size=13)

# ───────────────────────── 6. Раздел 02
section('02', 'Жизненный цикл заказа', 'Фаза заказа и цепочка согласования')

# ───────────────────────── 7. Фаза заказа
s = content('Фаза заказа — один системный статус',
            'Фазу задаёт ровно один системный код. Это проверено на всех 422 774 заказах выгрузки PM-06. '
            'Позиции без статуса встречаются только в плане 2027 года: это позиции графика ППР, по которым '
            'заказ ещё не создан.')
ph = [
    ('нет статуса', 'Позиция ППР', 'Заказа ещё нет. Не план и не потребность', LGREY, DARK),
    ('ОТКР', 'Открыт', 'Планирование и согласование', GREY, WHITE),
    ('ДЕБЛ · ЧДЕБ', 'Деблокирован', 'Разрешены проводки: подтверждения, списание МТР, счета', DARK, WHITE),
    ('ТЗКР', 'Технически закрыт', 'Только из ДЕБЛ. Работы завершены', GREEN, DARK),
    ('ЗАКР', 'Закрыт коммерчески', 'Только из ТЗКР при нулевом сальдо', GREEN, DARK),
]
pw = (W - 0.4) / 5
for i, (code, t, d, f, c) in enumerate(ph):
    x = X0 + i * (pw + 0.1)
    chevron(s, x, Y0 + 0.25, pw, 1.0, f, [{'t': code, 'bold': True, 'align': 'c'}],
            first=(i == 0), size=17, color=c)
    tb(s, x + 0.12, Y0 + 1.45, pw - 0.35, 1.5, [{'t': t, 'bold': True, 'size': 15, 'space': 6}, d],
       size=13)
box(s, X0, 4.35, 7.4, 1.95, fill=DARK, radius=0.18)
tb(s, X0 + 0.35, 4.5, 6.8, 1.7, [
    {'t': 'Заказ закрыт, если стоит ТЗКР или ЗАКР', 'bold': True, 'size': 22, 'color': GREEN, 'space': 8},
    {'t': 'Остаток «план − факт» у закрытого заказа — не потребность. Заказ, закрытый без факта, — '
          'это отменённая или перенесённая работа.', 'color': WHITE},
], size=14, anchor='m')
box(s, X0 + 7.7, 4.35, W - 7.7, 1.95, fill=WHITE, radius=0.18, anchor='m', paras=[
    {'t': 'Удалённые заказы (МТКУ) в отчёты BW не попадают', 'bullet': True},
    {'t': 'Позиции ППР — номера 1ХХХХХХХХХХХ из диапазона планов ТОРО', 'bullet': True},
], size=13, margin=(0.3, 0.1, 0.25, 0.1), space=8)

# ───────────────────────── 8. Согласование
s = content('Согласование заказа: кто ставит статус',
            'Порядковые пользовательские статусы — цепочка согласования; активен всегда один. '
            'ССПЛ и СГГС блокируют заказ для изменений. По упрощённой схеме заказ, укладывающийся в '
            'настроечные условия, деблокируют без СГГС. Аварийный заказ инженер ТОиР деблокирует сразу '
            'после создания.')
ap = [
    ('ПЛАН', '1', 'Система', 'Ставится автоматически при создании заказа'),
    ('СГПЛ', '2', 'Планировщик ТОиР', 'Проверил лимиты, присвоил СПП-элемент, отправил на согласование'),
    ('ССПЛ', '3', 'Старший планировщик', 'Проанализировал в ZPM_PLAN и согласовал. Заказ блокируется для изменений'),
    ('СГГС', '4', 'Главный специалист ТОиР', 'Проверил корректность планирования и согласовал'),
    ('ДЕБЛ', '→', 'Старший планировщик', 'Деблокирует согласованные заказы: можно отражать факт'),
]
aw = (W - 0.4) / 5
for i, (code, n, role, d) in enumerate(ap):
    x = X0 + i * (aw + 0.1)
    hl = code == 'СГПЛ'
    chevron(s, x, Y0 + 0.2, aw, 0.85, GREEN if hl else (DARK if i < 4 else GREY),
            [{'t': code, 'bold': True, 'align': 'c'}], first=(i == 0), size=18,
            color=DARK if hl else WHITE)
    box(s, x, Y0 + 1.25, aw - 0.12, 2.35, fill=TINT if hl else WHITE, radius=0.16)
    tb(s, x + 0.18, Y0 + 1.38, aw - 0.45, 2.15, [
        {'t': role, 'bold': True, 'size': 14, 'space': 6},
        {'t': d, 'color': MUTED},
    ], size=13)
tips = [
    ('Один активный', 'Порядковый статус всегда один. Два сразу — заказ был согласован и сброшен обратно в план'),
    ('Упрощённая схема', 'Заказ в пределах настроечных условий деблокируют без согласования старшим и главным'),
    ('Активный ЗНК', 'Пока есть необработанный запрос на корректировку, сроки и статусы согласования не меняются'),
]
tw = (W - 0.6) / 3
for i, (t, d) in enumerate(tips):
    x = X0 + i * (tw + 0.3)
    box(s, x, 5.05, tw, 1.35, fill=DARK, radius=0.16)
    tb(s, x + 0.22, 5.15, tw - 0.44, 1.2, [
        {'t': t, 'bold': True, 'color': GREEN, 'size': 14, 'space': 4},
        {'t': d, 'color': WHITE},
    ], size=12)

# ───────────────────────── 9. Статусы планов и бюджета
s = content('Статусы планов, бюджета и материалов',
            'Эти пользовательские статусы ставятся независимо друг от друга и фазу заказа не меняют. '
            'Они показывают, в какой план заказ включён и утверждён, что с бюджетом и материалами.')
cols = [
    ('Годовой план', GREEN, DARK, [('ГОД', 'включён в годовой план'), ('ПТОГ', 'согласован ПТО'),
                                   ('ГИП', 'согласован главным инженером'),
                                   ('УТВГ', 'утверждён в программе ремонтов'),
                                   ('УТВП', 'в программе и бюджете (авто из MCB)')]),
    ('Месяц и неделя', DARK, WHITE, [('МЕС', 'включён в месячный план'), ('ПТОМ', 'согласован ПТО'),
                                     ('УТВМ', 'утверждён в месячном плане'),
                                     ('НЕД', 'включён в недельный план'), ('УТВН', 'утверждён в недельном')]),
    ('Бюджет', DARK, WHITE, [('ТКБЕ', 'бюджет утверждён ТК БЕ'), ('ТРКБ', 'нужна корректировка'),
                             ('КОРБ', 'выделен и скорректирован'), ('ОТКБ', 'корректировка отклонена'),
                             ('ПЗТГ · ПЗТМ', 'прогноз завершения года / месяца')]),
    ('Материалы', GREY, WHITE, [('МТРЦ', 'МТР оценены'), ('МТРН', 'есть МТР без цены'),
                                ('СРОЧ', 'срочная закупка'), ('ТОПЗ · СОПЗ', 'опережающая закупка'),
                                ('НВСО', 'невостребованные остатки')]),
]
qw = (W - 0.75) / 4
for i, (t, hf, hc, rows) in enumerate(cols):
    x = X0 + i * (qw + 0.25)
    box(s, x, Y0 + 0.15, qw, 0.85, fill=hf, radius=0.16)
    box(s, x, Y0 + 0.65, qw, 4.55, fill=WHITE, radius=0.16)
    tb(s, x, Y0 + 0.15, qw, 0.5, [{'t': t, 'bold': True, 'align': 'c'}], size=15, color=hc, anchor='m')
    for j, (code, d) in enumerate(rows):
        yy = Y0 + 0.85 + j * 0.86
        box(s, x + 0.14, yy, 1.08, 0.42, fill=LGREY if code != 'УТВГ' else GREEN, radius=0.21,
            paras=[{'t': code, 'bold': True, 'align': 'c'}], size=10, anchor='m', margin=(0.01, 0, 0.01, 0))
        tb(s, x + 1.3, yy - 0.04, qw - 1.4, 0.8, [d], size=11, anchor='t')

# ───────────────────────── 10. Раздел 03
section('03', 'Выполнение и закрытие', 'Цепочка закрытия АО «Развитие» и контроль в IW38')

# ───────────────────────── 11. Цепочка закрытия
s = content('Цепочка закрытия заказа в АО «Развитие»',
            'Для работ «Развития» добавлен шаг приёмки службой заказчика БЕ: ПРСЗ или НПСЗ. ВСБЕ ставится '
            'автоматически при регистрации выручки в ZMM_370. По памятке ТОиР-84 ФХСМ означает, что затрат '
            'больше не ожидается и заказ готов к ТЗКР.')
chain = [
    ('ДЕБЛ', 'Деблокирован', 'Проводки разрешены', DARK, WHITE),
    ('ПДТВ · ЧПДТ', 'Подтверждения', 'Трудозатраты ХС подтверждены полностью или частично', GREY, WHITE),
    ('ПРНТ', 'Работы приняты', 'Для заказов с УСО. НПРН — не приняты', GREY, WHITE),
    ('ФХСМ', 'Факт проведён', 'ХС, работы подрядчика, МТР — затрат больше не будет', GREY, WHITE),
    ('ТЗКР', 'Тех. закрытие', 'Работы завершены, заказ закрыт технически', GREEN, DARK),
    ('ПРСЗ · НПСЗ', 'Приёмка СЗ', 'Служба заказчика БЕ приняла или не приняла работы', DARK, WHITE),
    ('ВСБЕ', 'Выставлен в БЕ', 'Автоматически при регистрации выручки (ZMM_370)', DARK, WHITE),
    ('ЗАКР', 'Закрыт', 'Коммерческое закрытие, сальдо = 0', GREEN, DARK),
]
cw4 = (W - 0.3) / 4
for i, (code, t, d, f, c) in enumerate(chain):
    row, col = divmod(i, 4)
    x = X0 + col * (cw4 + 0.1)
    y = Y0 + 0.15 + row * 2.6
    chevron(s, x, y, cw4, 0.9, f, [{'t': code, 'bold': True, 'align': 'c'}],
            first=(col == 0), size=17, color=c)
    tb(s, x + 0.15, y + 1.05, cw4 - 0.45, 1.4, [{'t': t, 'bold': True, 'size': 15, 'space': 5},
                                                 {'t': d, 'color': MUTED}], size=13)
box(s, X0 + 0.0, Y0 + 2.45, W, 0.02, fill=LGREY)

# ───────────────────────── 12. Проверки
s = content('Что проверяет система на каждом шаге',
            'Проверки из памятки ТОиР-77 для транзакций IW32 и IW38. Статусы ПРНТ/НПРН и ПРСЗ/НПСЗ '
            'взаимоисключающие: при установке одного система сама снимает другой.')
rows = [
    ('Поставить ПРНТ', 'Заказ деблокирован (ДЕБЛ)', 'Статус не ставится'),
    ('Технически закрыть (ТЗКР)', 'Есть ПРНТ — для заказов с УСО, или ПДТВ/ЧПДТ — для заказов с ФОТ. Нет НПРН', 'Закрыть нельзя'),
    ('Поставить ПРСЗ или НПСЗ', 'Заказ в ТЗКР, у сотрудника роль «Специалист службы заказчика БЕ»', 'Статус не ставится'),
    ('После ПРСЗ', 'Запрещены счета (RMRP), подтверждения (RMRU, RMPH), списание МТР (RMWA), акты подрядчика (RMWE)', 'Проводка отклоняется'),
    ('Снять ПРСЗ', 'Только роль «Руководитель службы заказчика БЕ» и только до ВСБЕ', 'После ВСБЕ — ошибка'),
]
cxs = [X0, X0 + 3.1, X0 + 9.55]
cws = [3.0, 6.35, W - 9.55 + 0.0]
box(s, X0, Y0 + 0.1, W, 0.55, fill=DARK, radius=0.27)
for j, h in enumerate(['Действие', 'Условие', 'Если нарушено']):
    tb(s, cxs[j] + 0.25, Y0 + 0.1, cws[j] - 0.3, 0.55, [{'t': h, 'bold': True}], size=13,
       color=GREEN if j == 0 else WHITE, anchor='m')
rh = 0.72
for i, (a, b, c) in enumerate(rows):
    y = Y0 + 0.78 + i * (rh + 0.08)
    box(s, X0, y, W, rh, fill=WHITE, radius=0.14)
    tb(s, cxs[0] + 0.25, y, cws[0] - 0.3, rh, [{'t': a, 'bold': True}], size=13, anchor='m')
    tb(s, cxs[1] + 0.25, y, cws[1] - 0.3, rh, [b], size=12, anchor='m')
    box(s, cxs[2] + 0.2, y + 0.16, cws[2] - 0.4, 0.4, fill=LGREY, radius=0.2,
        paras=[{'t': c, 'align': 'c'}], size=11, anchor='m', margin=(0.05, 0, 0.05, 0))
tb(s, X0, 5.95, W, 0.5, [{'t': [B('При сохранении — предупреждения: '),
                                ('МТР подрядчика не сходятся с услугой по плану или факту; при ТЗКР счета по УСО проведены не полностью; '
                                 'ОС заказа не совпадает с ОС единицы оборудования; дата списания ОС раньше базисного начала.', {})]}],
   size=12, color=MUTED)

# ───────────────────────── 13. IW38 контроль
s = content('Контроль закрытия в IW38',
            'Блок «Параметры контроля закрытия заказов» в IW38 (памятка ТОиР-84). Техник по учёту добивается, '
            'чтобы все заказы были готовы к ТЗКР. Отчёт CONTROL — минимум раз в неделю для инженеров ТОиР и '
            'руководителей: базисное окончание — вчерашняя дата, проблемные заказы подсвечены красным. '
            'При большом объёме данных отключайте «Дополнительные поля».')
lw = 7.3
box(s, X0, Y0 + 0.1, lw, 4.1, fill=WHITE, radius=0.18)
tb(s, X0 + 0.3, Y0 + 0.22, lw - 0.6, 0.45, [{'t': 'Статус отражения факта', 'bold': True}], size=16)
tb(s, X0 + 0.3, Y0 + 0.62, lw - 0.6, 0.4, ['Код = полнота + категория затрат, например 1Н-СМ — списание МТР не проведено'],
   size=12, color=MUTED)
grp = [('1Н', 'не проведено', DARK, WHITE), ('2Ч', 'частично', GREY, WHITE), ('3П', 'проведено', GREEN, DARK)]
for i, (c, t, f, fc) in enumerate(grp):
    box(s, X0 + 0.3 + i * 1.3, Y0 + 1.15, 1.2, 0.75, fill=f, radius=0.14, anchor='m',
        paras=[{'t': c, 'bold': True, 'align': 'c', 'size': 16, 'space': 0},
               {'t': t, 'align': 'c', 'size': 10}], color=fc, margin=(0.02, 0.02, 0.02, 0.02))
cats = [('ХС', 'подтверждение работ ХС'), ('РП', 'подтверждение работ подрядчика'),
        ('СМ', 'списание МТР'), ('СФ', 'счёт-фактура по УСО'), ('СФП', 'счёт-фактура по МТР подрядчика')]
for j, (c, t) in enumerate(cats):
    yy = Y0 + 2.12 + j * 0.4
    box(s, X0 + 0.3, yy, 0.8, 0.32, fill=LGREY, radius=0.16,
        paras=[{'t': c, 'bold': True, 'align': 'c'}], size=11, anchor='m', margin=(0, 0, 0, 0))
    tb(s, X0 + 1.25, yy - 0.02, 5.5, 0.36, [t], size=12, anchor='m')
tb(s, X0 + 4.3, Y0 + 1.15, 2.75, 0.9, [{'t': '× 5 категорий', 'bold': True, 'size': 14, 'space': 2},
                                       {'t': 'на каждый уровень полноты', 'color': MUTED}], size=11, anchor='m')
rx = X0 + lw + 0.3
rw = W - lw - 0.3
box(s, rx, Y0 + 0.1, rw, 1.25, fill=WHITE, radius=0.18)
tb(s, rx + 0.25, Y0 + 0.2, rw - 0.5, 1.1, [
    {'t': 'Статус периода проводок', 'bold': True, 'size': 15, 'space': 4},
    'Проводки вне базисных сроков заказа — по каждой категории затрат'], size=12)
box(s, rx, Y0 + 1.55, rw, 2.65, fill=DARK, radius=0.18)
tb(s, rx + 0.25, Y0 + 1.65, rw - 0.5, 0.45, [{'t': 'Статус приёмки', 'bold': True}], size=15, color=GREEN)
acc = [('#ТЗКР', 'всё проведено — пора закрыть'), ('#ПРСЗ', 'пора принять службой заказчика'),
       ('#НПСЗ', 'заказ блокирует выручку'), ('#ВСБЕ', 'выставлен вне периода приёмки')]
for j, (c, t) in enumerate(acc):
    yy = Y0 + 2.15 + j * 0.49
    box(s, rx + 0.25, yy, 1.0, 0.36, fill=GREEN, radius=0.18,
        paras=[{'t': c, 'bold': True, 'align': 'c'}], size=11, anchor='m', margin=(0, 0, 0, 0))
    tb(s, rx + 1.4, yy - 0.02, rw - 1.6, 0.4, [t], size=12, color=WHITE, anchor='m')
box(s, X0, 5.5, W, 0.85, fill=TINT, radius=0.16, anchor='m', paras=[
    {'t': [B('Раз в неделю: '), ('IW38 → формат «CONTROL — контроль заказов» → своё ТМ, базисное окончание = вчера → '
                                  'фильтры «Нет подтверждения работ» и «Нет списания МТР». Отменили работу — закройте заказ '
                                  'и пометьте на удаление или перенесите сроки.', {})]}],
    size=13, margin=(0.25, 0.05, 0.25, 0.05))

# ───────────────────────── 14. Раздел 04
section('04', 'Материалы', 'Когда закупка увидит потребность заказа')

# ───────────────────────── 15. Резерв./заявка
s = content('«Резерв./заявка»: когда закупка видит потребность',
            'Признак релевантности компонента для планирования потребности в материалах (ППМ). Пример из '
            'выгрузки PM-06 на 14.09.2026 по экскаваторам WK: на 2027 год не покрыто 525 млн ₽ потребности, '
            'из них 393 млн ₽ — на строках «Начиная с деблокирования» в ещё не деблокированных заказах. '
            'Закупка их пока не видит.')
kw3 = 2.75
opts = [
    ('Немедленно', GREEN, DARK, 'Потребность уходит в ППМ сразу при сохранении заказа. Значение по умолчанию'),
    ('Начиная с деблокирования', GREY, WHITE, 'В ППМ только после ДЕБЛ. До деблокирования закупка потребность не видит'),
    ('Никогда', DARK, WHITE, 'В ППМ и закупочную заявку не попадает. Покрыть можно только со склада'),
]
for i, (t, hf, hc, d) in enumerate(opts):
    x = X0 + i * (kw3 + 0.25)
    card(s, x, Y0 + 0.15, kw3, 3.3, t, [d], hfill=hf, hcolor=hc, hsize=14, bsize=14, hh=0.75)
tb(s, X0, Y0 + 3.7, 3 * kw3 + 0.5, 1.4, [
    {'t': [B('«Никогда» ставится автоматически '), ('справочным позициям (S) вместе с «Калькуляция: не релевантно» '
                                                  'и движениям оборотного фонда с неисправной партией.', {})], 'bullet': True},
    {'t': [B('Поздний деблок = поздняя закупка. '), ('Планируйте деблокирование с запасом на срок поставки МТР.', {})],
     'bullet': True},
], size=13, space=8)
sx = X0 + 3 * kw3 + 0.75
sw2 = W - (sx - X0)
box(s, sx, Y0 + 0.15, sw2, 5.1, fill=DARK, radius=0.2)
tb(s, sx + 0.3, Y0 + 0.4, sw2 - 0.6, 4.7, [
    {'t': 'Пример: WK, план 2027', 'size': 13, 'color': LGREY, 'space': 10},
    {'t': '393 млн ₽', 'bold': True, 'size': 30, 'color': GREEN, 'space': 2},
    {'t': 'из 525 млн ₽ дефицита — строки «Начиная с деблокирования»: закупка их пока не видит',
     'color': WHITE, 'size': 13, 'space': 16},
    {'t': '81 млн ₽', 'bold': True, 'size': 24, 'color': WHITE, 'space': 2},
    {'t': 'строки «Никогда»: в заявку не попадут совсем', 'color': LGREY, 'size': 13},
], size=13)
source(s, 'Источник: выгрузка BW PM-M06 на 14.09.2026, расчёт обеспеченности WK')

# ───────────────────────── 16. СРОЧ и НВСО
s = content('Срочность (СРОЧ) и невостребованные остатки (НВСО)',
            'Памятки ТОиР-75, ТОиР-86 и ТОиР-96. Проверки срабатывают при изменении компонентов и сохранении '
            'заказа и берут данные ППМ из ночного среза (offline). Расхождения за день выравнивает ночной '
            'фоновый пересчёт; вручную — функция «ФОНОВЫЙ пересчёт заказов ТОРО» в IW38.')
hw = (W - 0.3) / 2
for k, (t, hf, hc) in enumerate([('СРОЧ — срочная потребность', GREEN, DARK),
                                 ('НВСО — невостребованные остатки', DARK, WHITE)]):
    x = X0 + k * (hw + 0.3)
    box(s, x, Y0 + 0.1, hw, 0.9, fill=hf, radius=0.16)
    box(s, x, Y0 + 0.6, hw, 4.65, fill=WHITE, radius=0.16)
    tb(s, x, Y0 + 0.1, hw, 0.5, [{'t': t, 'bold': True, 'align': 'c'}], size=16, color=hc, anchor='m')
x = X0
box(s, x + 0.3, Y0 + 0.85, hw - 0.6, 1.25, fill=TINT, radius=0.14, anchor='m', paras=[
    {'t': 'Позиция срочная, если', 'bold': True, 'space': 4, 'size': 12, 'color': MUTED},
    {'t': 'сегодня + срок поставки + срок согласования > дата потребности', 'bold': True, 'space': 4},
    {'t': 'и свободных остатков не хватает', 'bold': True},
], size=14, margin=(0.2, 0.05, 0.2, 0.05))
tb(s, x + 0.3, Y0 + 2.3, hw - 0.6, 2.9, [
    {'t': 'Остатки считаются с перевалочной базой и прямыми аналогами МТО, без запасов проектов', 'bullet': True},
    {'t': [B('Не передана в ППМ: '), ('признак снимается, когда остатков стало хватать', {})], 'bullet': True},
    {'t': [B('Передана в ППМ («Немедленно»): '), ('снимается только по сроку — когда дата вышла за срок поставки', {})],
     'bullet': True},
], size=13, space=8)
x = X0 + hw + 0.3
tb(s, x + 0.3, Y0 + 0.85, hw - 0.6, 2.7, [
    {'t': 'При удалении, сокращении или переносе потребности система сравнивает:', 'space': 6},
    {'t': 'заявленную потребность на 2 года', 'bullet': True, 'space': 2},
    {'t': 'запас на складе + спецификации к договору', 'bullet': True, 'space': 8},
    {'t': [B('Профицит в текущем году — '), ('предупреждение', {})], 'bullet': True, 'bcolor': GREY, 'space': 2},
    {'t': [B('Профицит в следующем году — '), ('ошибка: удалить может главный специалист ТОиР УК', {})],
     'bullet': True, 'bcolor': GREY},
], size=13, space=4)
box(s, x + 0.3, Y0 + 3.7, hw - 0.6, 1.4, fill=DARK, radius=0.14, anchor='m', paras=[
    {'t': 'Перед сокращением проверьте', 'bold': True, 'color': GREEN, 'space': 4},
    {'t': 'ZMM_MRP_MAINT: варианты NVSO_SVERKA и NVSO_ZNPSKL, склад — MB52. После удаления или ТЗКР '
          'с отказом ставится НВСО, его не снять.', 'color': WHITE},
], size=12, margin=(0.25, 0.05, 0.25, 0.05))

# ───────────────────────── 17. Раздел 05
section('05', 'Перенос и отмена работ', 'Копии заказов, блокировки сроков, запросы на корректировку')

# ───────────────────────── 18. Перенос в «Развитие»
s = content('Перенос в «Развитие» — это копия, а не переназначение',
            'Заказы производственной БЕ мигрируют в «Развитие» копированием вместе со связанными сообщениями: '
            'функция «Массовое копирование заказов/сообщений» в IW38/IW39. Связь сообщений хранится в IW28/IW29 '
            '(поля «Скопированное сообщение», «Ссылка клиента»). По выгрузкам PM-06 найдено 1 725 пар '
            '«та же ЕО + вид работ + дата начала», в 147 из них факта нет ни у одной стороны.')
bw_ = 3.6
box(s, X0, Y0 + 0.3, bw_, 1.7, fill=WHITE, radius=0.18, line=LGREY)
tb(s, X0 + 0.25, Y0 + 0.4, bw_ - 0.5, 1.5, [
    {'t': 'Заказ БЕ', 'bold': True, 'size': 17, 'space': 4},
    {'t': 'Оригинал на планирующем заводе производственной БЕ', 'color': MUTED}], size=13, anchor='m')
chevron(s, X0 + bw_ + 0.2, Y0 + 0.65, W - 2 * bw_ - 0.4, 1.0, LGREY,
        [{'t': 'Массовое копирование', 'bold': True, 'align': 'c', 'space': 0}, {'t': 'заказов/сообщений в IW38', 'align': 'c'}], first=True, size=13)
box(s, X1 - bw_, Y0 + 0.3, bw_, 1.7, fill=DARK, radius=0.18)
tb(s, X1 - bw_ + 0.25, Y0 + 0.4, bw_ - 0.5, 1.5, [
    {'t': 'Заказ «Развития»', 'bold': True, 'size': 17, 'space': 4, 'color': GREEN},
    {'t': 'Копия с заводом филиала и копиями сообщений', 'color': WHITE}], size=13, anchor='m')
rules = [
    ('Переназначить нельзя', 'Заказ, созданный не на тот завод, удаляют и создают заново. В IW31 и IW34 завод выбирают на первом экране'),
    ('Сообщение с заказом', 'Отвязать заказ → переназначить сообщение между БЕ → удалить ошибочный заказ → получатель создаёт новый'),
    ('Как считать', 'Факт — с обеих сторон. Неисполненный план — только по копии «Развития», иначе он задваивается'),
]
rw3 = (W - 0.6) / 3
for i, (t, d) in enumerate(rules):
    x = X0 + i * (rw3 + 0.3)
    card(s, x, 3.55, rw3, 2.55, t, [d], hfill=GREEN if i == 2 else DARK, hcolor=DARK if i == 2 else WHITE,
         bsize=15, hsize=16)

# ───────────────────────── 19. Отмена, блокировки, ЗНК
s = content('Отмена, перенос сроков и блокировки',
            'Отменённую или перенесённую работу закрывают и (или) помечают на удаление. Помеченные на удаление '
            'заказы в выгрузку не попадают; закрытые без факта видны как ТЗКР/ЗАКР с нулевым фактом — по WK '
            'за 2026 год таких 67. Связи старого и нового заказа при переносе в PM-06 нет.')
blocks = [
    ('Работу отменили или перенесли', GREEN, DARK, [
        {'t': 'Закрыть заказ и пометить на удаление — или перенести сроки', 'bullet': True},
        {'t': 'Не оставлять деблокированный заказ без факта: он искажает исполнение программы', 'bullet': True},
        {'t': 'Удаляя законтрактованные МТР, проверить НВСО', 'bullet': True},
    ]),
    ('Блокировка сроков · ZPM_DATE_BLOCK', DARK, WHITE, [
        {'t': 'Задаётся по БЕ, заводу, группе плановиков, ТМ и ЕО на период', 'bullet': True},
        {'t': 'Заказ попадает под блокировку по базисной дате начала', 'bullet': True},
        {'t': 'Запрещает менять сроки и простой, а также закрывать заказ', 'bullet': True},
    ]),
    ('Активный ЗНК к заказу', DARK, WHITE, [
        {'t': 'Сроки и статусы согласования (СГПЛ, ССПЛ, СГГС) не меняются', 'bullet': True},
        {'t': 'ZMM_CHREQ_MAINT: номер заказа с двумя ведущими нулями, поле «Автор ЗНК» очистить', 'bullet': True},
        {'t': 'Каждый ЗНК согласовать полностью или отменить', 'bullet': True},
    ]),
]
bw3 = (W - 0.6) / 3
for i, (t, hf, hc, body) in enumerate(blocks):
    card(s, X0 + i * (bw3 + 0.3), Y0 + 0.15, bw3, 3.75, t, body, hfill=hf, hcolor=hc, bsize=14, hsize=15,
         hh=0.7, space=10)
box(s, X0, 5.3, W, 1.05, fill=DARK, radius=0.18)
tb(s, X0 + 0.35, 5.3, 2.4, 1.05, [{'t': '67 заказов', 'bold': True}], size=26, color=GREEN, anchor='m')
tb(s, X0 + 2.8, 5.3, W - 3.1, 1.05, [
    'WK за 2026 год закрыты без факта: это отменённые или перенесённые работы. Связи старого и нового '
    'заказа при переносе в PM-06 нет, поэтому новый заказ ищут по ЕО и виду работ.'], size=13, color=WHITE, anchor='m')

# ───────────────────────── 20. Раздел 06
section('06', 'Практика', 'Чтение PM-06, типичные ошибки, чек-лист, самопроверка')

# ───────────────────────── 21. Чтение PM-06 + диаграмма
s = content('Как читать PM-06: состояние заказа за 5 шагов',
            'Правила подтверждены на выгрузках 2022–2027 годов. Диаграмма — программа ремонтов WK 2026 года '
            'по стадиям SAP на 14.09.2026: закрыто (ТЗКР и далее) 52 % плана в рублях; 93 деблокированных заказа '
            'на 601 млн ₽ не имеют ни факта, ни подтверждений.')
steps5 = [
    ('Нет статусов', 'позиция ППР — не заказ, в план и потребность не включать'),
    ('ТЗКР или ЗАКР', 'заказ закрыт: остаток «план − факт» не потребность'),
    ('ДЕБЛ', 'в работе: смотрим ПДТВ/ЧПДТ, ПРНТ, ФХСМ и факт'),
    ('ОТКР', 'на согласовании (ПЛАН–ССПЛ) или согласован (СГГС)'),
    ('Пара БЕ + «Развитие»', 'та же ЕО, вид работ и дата начала — план только по копии'),
]
lw5 = 5.35
for i, (c, d) in enumerate(steps5):
    y = Y0 + 0.1 + i * 1.04
    box(s, X0, y, lw5, 0.92, fill=WHITE, radius=0.16)
    box(s, X0 + 0.18, y + 0.19, 0.54, 0.54, fill=GREEN if i == 1 else DARK, radius=0.27,
        paras=[{'t': str(i + 1), 'bold': True, 'align': 'c'}], size=15,
        color=DARK if i == 1 else GREEN, anchor='m', margin=(0, 0, 0, 0))
    tb(s, X0 + 0.9, y + 0.06, lw5 - 1.05, 0.82, [{'t': c, 'bold': True, 'size': 14, 'space': 1},
                                                  {'t': d, 'color': MUTED}], size=12, anchor='m')
cx = X0 + lw5 + 0.3
cwid = W - lw5 - 0.3
box(s, cx, Y0 + 0.1, cwid, 5.12, fill=WHITE, radius=0.18)
tb(s, cx + 0.3, Y0 + 0.2, cwid - 0.6, 0.4, [{'t': 'WK, программа 2026: план по стадиям, млн ₽', 'bold': True}], size=14)
stages = [('На согласовании', 503.3), ('Согласован (СГГС)', 610.7), ('Деблокирован, без факта', 601.2),
          ('В работе', 146.0), ('Факт проведён (ФХСМ)', 28.5), ('Тех. закрыт (ТЗКР)', 37.6),
          ('Принят СЗ (ПРСЗ)', 135.5), ('Выставлен в БЕ (ВСБЕ)', 729.3), ('Закрыт (ЗАКР)', 1159.6)]
cd = CategoryChartData()
cd.categories = [a for a, _ in reversed(stages)]
cd.add_series('План, млн ₽', [v for _, v in reversed(stages)])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(cx + 0.15), Inches(Y0 + 0.6),
                        Inches(cwid - 0.35), Inches(4.3), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = False
pl = ch.plots[0]
pl.gap_width = 45
pl.has_data_labels = True
dl = pl.data_labels
dl.number_format = '# ##0'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(11)
dl.font.name = FONT
dl.font.color.rgb = rgb(DARK)
ser = pl.series[0]
closed_n = 4  # последние четыре стадии (ТЗКР и далее) — закрыто
for idx in range(len(stages)):
    pt = ser.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(GREEN if idx < closed_n else GREY)
ca = ch.category_axis
ca.tick_labels.font.size = Pt(11)
ca.tick_labels.font.name = FONT
ca.tick_labels.font.color.rgb = rgb(DARK)
ca.format.line.color.rgb = rgb(LGREY)
ca.has_major_gridlines = False
va = ch.value_axis
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = 1450
tb(s, cx + 0.3, Y0 + 4.85, cwid - 0.6, 0.3, [{'t': [('■ ', {'color': GREEN, 'bold': True}), ('закрыто: ТЗКР и далее    ', {}),
                                                   ('■ ', {'color': GREY, 'bold': True}), ('ещё не закрыто', {})]}],
   size=11, color=MUTED)

# ───────────────────────── 22. Типичные ошибки
s = content('Типичные ошибки и как их избежать',
            'Ошибки собраны по проверкам SAP из памяток и по тому, что видно в выгрузках PM-06.')
errs = [
    ('Заказ не на том заводе', 'Переназначить нельзя', 'Удалить и создать заново в IW31/IW34 с заводом того, чьи ресурсы работают'),
    ('Деблокирован, но пусто', 'Нет факта и подтверждений — исполнение искажено', 'Еженедельный CONTROL в IW38; отменили — закрыть и пометить на удаление'),
    ('Поздний деблок', '«Начиная с деблокирования»: закупка не видит потребность', 'Деблокировать с запасом на срок поставки МТР'),
    ('Не даёт ТЗКР', 'Нет ПРНТ или подтверждений, либо стоит НПРН', 'Проверить «Статус отражения факта» в IW38 и довести факт'),
    ('Сокращение МТР → ошибка', 'Законтрактованная потребность даст профицит', 'До сокращения — ZMM_MRP_MAINT (NVSO_SVERKA, NVSO_ZNPSKL)'),
    ('Оригинал после копии', 'План задваивается в отчётах', 'Проверять пары по ссылкам сообщений в IW28 и считать план по копии'),
]
ew, eh = (W - 0.6) / 3, 2.45
for i, (t, cons, fix) in enumerate(errs):
    x = X0 + (i % 3) * (ew + 0.3)
    y = Y0 + 0.1 + (i // 3) * (eh + 0.25)
    box(s, x, y, ew, eh, fill=WHITE, radius=0.18)
    box(s, x + 0.2, y + 0.2, 0.46, 0.46, fill=DARK, radius=0.23,
        paras=[{'t': str(i + 1), 'bold': True, 'align': 'c'}], size=13, color=GREEN, anchor='m', margin=(0, 0, 0, 0))
    tb(s, x + 0.8, y + 0.17, ew - 1.0, 0.52, [{'t': t, 'bold': True}], size=15, anchor='m')
    tb(s, x + 0.2, y + 0.82, ew - 0.4, 0.6, [cons], size=12, color=MUTED)
    box(s, x + 0.2, y + 1.42, ew - 0.4, 0.88, fill=TINT, radius=0.12, anchor='m',
        paras=[{'t': fix}], size=12, margin=(0.14, 0.04, 0.14, 0.04))

# ───────────────────────── 23. Чек-лист
s = content('Чек-лист планировщика',
            'Ритм работы, который закрывает типичные ошибки. Недельные пункты — к планёрке, месячные — к '
            'согласованию план-графика и закрытию периода.')
lists = [
    ('Каждую неделю', GREEN, DARK, [
        'IW38 CONTROL: нет подтверждений, нет списаний МТР',
        'Деблокированные заказы без факта: провести, перенести или закрыть',
        'Заказы следующей недели отправлены на согласование (СГПЛ)',
        'Активные ЗНК к своим заказам обработаны',
        'Срочные позиции (СРОЧ): есть ли покрытие к дате работ',
    ]),
    ('Каждый месяц', DARK, WHITE, [
        'Заказы с ФХСМ технически закрыты (ТЗКР)',
        'Статус приёмки: #ТЗКР и #ПРСЗ отработаны',
        'Прогноз завершения месяца (ПЗТМ) согласован',
        'План-график следующего месяца: МЕС → УТВМ',
        'Удаления МТР не создали лишних НВСО',
    ]),
    ('Перед годом', GREY, WHITE, [
        'Годовой план: ГОД → ПТОГ → ГИП → УТВГ',
        'ППР распределены по правилу ресурсов: обходы — «Развитие», ЕТО — БЕ',
        'Ответственные в ТМ, ЕО и ППР проверены до отзыва',
        'Деблокирование запланировано с учётом сроков поставки',
    ]),
]
lw3 = (W - 0.6) / 3
for i, (t, hf, hc, items) in enumerate(lists):
    x = X0 + i * (lw3 + 0.3)
    box(s, x, Y0 + 0.1, lw3, 0.9, fill=hf, radius=0.16)
    box(s, x, Y0 + 0.6, lw3, 4.65, fill=WHITE, radius=0.16)
    tb(s, x, Y0 + 0.1, lw3, 0.5, [{'t': t, 'bold': True, 'align': 'c'}], size=16, color=hc, anchor='m')
    for j, it in enumerate(items):
        yy = Y0 + 0.85 + j * 0.86
        box(s, x + 0.22, yy + 0.06, 0.3, 0.3, fill=WHITE, line=GREEN if i == 0 else GREY, radius=0.06, lw=1.5)
        tb(s, x + 0.65, yy, lw3 - 0.85, 0.82, [it], size=13)

# ───────────────────────── 24. Самопроверка
s = content('Самопроверка',
            'Ответы — на следующем слайде. Предложите слушателям ответить самостоятельно, прежде чем показывать ответы.')
qs = [
    'Заказ в ТЗКР, факт по нему нулевой. Это потребность в МТР?',
    'Какой статус ставит старший планировщик и что он делает с заказом?',
    'Заказ согласован (СГГС), а закупка не видит его МТР. Почему?',
    'Заказ создан на завод БЕ, а работы выполняет «Развитие». Что делать?',
    'После какого статуса на заказ уже нельзя списать МТР?',
    'Почему не удаётся изменить сроки согласованного заказа?',
]
qw2 = (W - 0.3) / 2
for i, q in enumerate(qs):
    x = X0 + (i % 2) * (qw2 + 0.3)
    y = Y0 + 0.15 + (i // 2) * 1.72
    box(s, x, y, qw2, 1.5, fill=WHITE, radius=0.18)
    box(s, x + 0.25, y + 0.42, 0.66, 0.66, fill=GREEN, radius=0.33,
        paras=[{'t': str(i + 1), 'bold': True, 'align': 'c'}], size=18, anchor='m', margin=(0, 0, 0, 0))
    tb(s, x + 1.15, y + 0.1, qw2 - 1.4, 1.3, [q], size=16, anchor='m')

# ───────────────────────── 25. Ответы
s = content('Ответы')
ans = [
    ('Нет', 'Заказ закрыт: работу выполнили без факта или отменили. Остаток «план − факт» не потребность'),
    ('ССПЛ', 'Согласовано старшим планировщиком — заказ блокируется для изменений'),
    ('«Начиная с деблокирования»', 'Потребность уйдёт в ППМ только после ДЕБЛ — нужно деблокировать заказ'),
    ('Пересоздать', 'Переназначить нельзя: удалить и создать с заводом филиала «Развития» или перенести копированием в IW38'),
    ('ПРСЗ', 'После приёмки службой заказчика запрещены списание МТР, подтверждения, счета и акты'),
    ('Блокировка', 'Активный ЗНК, блокировка сроков ZPM_DATE_BLOCK или статус ССПЛ/СГГС, блокирующий изменения'),
]
for i, (a, d) in enumerate(ans):
    x = X0 + (i % 2) * (qw2 + 0.3)
    y = Y0 + 0.15 + (i // 2) * 1.72
    box(s, x, y, qw2, 1.5, fill=WHITE, radius=0.18)
    box(s, x + 0.25, y + 0.42, 0.66, 0.66, fill=DARK, radius=0.33,
        paras=[{'t': str(i + 1), 'bold': True, 'align': 'c'}], size=18, color=GREEN, anchor='m', margin=(0, 0, 0, 0))
    tb(s, x + 1.15, y + 0.1, qw2 - 1.4, 1.3, [{'t': a, 'bold': True, 'size': 16, 'space': 3},
                                              {'t': d, 'color': MUTED}], size=13, anchor='m')

# ───────────────────────── 26. Справочник транзакций
s = content('Справочник транзакций',
            'Транзакции, упомянутые в методичке, с назначением для планировщика.')
tx = [
    ('IW31 · IW34', 'Создать заказ (без сообщения / к сообщению), выбор планирующего завода'),
    ('IW32', 'Изменить заказ: статусы, деблокирование'),
    ('IW38 · IW39', 'Список заказов: массовые статусы, копирование, контроль закрытия, фоновый пересчёт'),
    ('IW28 · IW29', 'Список сообщений: переназначение между БЕ, отвязка заказов'),
    ('ZPM_PLAN', 'АРМ планировщика: план-график, согласование'),
    ('IP17', 'Позиции ППР: массовая смена ответственных'),
    ('IL05 · IE05', 'Массовая смена ответственных в ТМ и ЕО'),
    ('ZPM_DATE_BLOCK', 'Блокировки сроков заказов'),
    ('ZMM_CHREQ_MAINT', 'Запросы на корректировку (ЗНК)'),
    ('ZMM_MRP_MAINT', 'Срез ППМ (offline), сверка НВСО'),
    ('MD04 · MB52', 'Ситуация ППМ онлайн · складские запасы'),
    ('ZPM_MTR_UPD', 'Массовое изменение и удаление МТР в заказах'),
    ('ZPM_MTRP', 'Массовая загрузка МТР подрядчика'),
    ('ZMM_370', 'Пульт по выручке: ставит ВСБЕ'),
]
half = (len(tx) + 1) // 2
tw2 = (W - 0.3) / 2
for i, (c, d) in enumerate(tx):
    col, row = divmod(i, half)
    x = X0 + col * (tw2 + 0.3)
    y = Y0 + 0.1 + row * 0.73
    box(s, x, y, tw2, 0.63, fill=WHITE, radius=0.14)
    box(s, x + 0.12, y + 0.11, 2.05, 0.41, fill=DARK, radius=0.2,
        paras=[{'t': c, 'bold': True, 'align': 'c'}], size=11, color=GREEN, anchor='m', margin=(0.02, 0, 0.02, 0))
    tb(s, x + 2.35, y, tw2 - 2.5, 0.63, [d], size=12, anchor='m')

# ───────────────────────── 27. Финал
s = prs.slides.add_slide(prs.slide_layouts[L_END])
for ph in s.placeholders:
    i = ph.placeholder_format.idx
    if i == 13:
        ph.text_frame.text = 'Вопросы и предложения'
    elif i == 14:
        ph.text_frame.text = 'SAP ТОРО для планировщика · 2026'
        _nobullet(ph.text_frame.paragraphs[0])
s.notes_slide.notes_text_frame.text = (
    'Источники: КД ERP.13.9.2.51–57 и .7; регламенты годового и оперативного планирования; памятки '
    'ТОиР-72 (блокировки), ТОиР-75 (законтрактованные потребности), ТОиР-77 (ключевые отличия SAP ТОРО '
    'в АО «Развитие»), ТОиР-79 (перенос заказов), ТОиР-84 (контроль закрытия в IW38), ТОиР-86 (срочность), '
    'ТОиР-91 (ЗНК), ТОиР-96 (СРОЧ/НВСО); памятка по контролю отражения факта; выгрузки BW PM-M06 2022–2027.')

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
