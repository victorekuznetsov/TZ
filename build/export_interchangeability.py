#!/usr/bin/env python3
"""Export the row-by-row interchangeability register for review."""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "exports"
HEAD = ["Лист", "Строка", "Тип связи", "Книги", "Номера деталей", "Отметки источника", "Комментарий"]


def rows():
    data = json.loads((ROOT / "data" / "interchange.json").read_text(encoding="utf-8"))
    for e in data["evidence"]:
        parts = e["parts"]
        yield [
            e["sheet"], e["row"],
            "Одинаковый номер" if e.get("relation") == "identity" else "Прямая замена по строке",
            " ↔ ".join(p["book"] for p in parts),
            " ↔ ".join(p["num"] for p in parts),
            "; ".join(e.get("flags") or []), e.get("note") or "",
        ]


def main():
    OUT.mkdir(exist_ok=True)
    body = list(rows())
    with (OUT / "interchangeability.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(HEAD)
        w.writerows(body)

    wb = Workbook()
    ws = wb.active
    ws.title = "Взаимозаменяемость"
    ws.append(HEAD)
    for r in body:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="166534")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col, width in {"A":12,"B":10,"C":23,"D":30,"E":58,"F":28,"G":70}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(OUT / "interchangeability.xlsx")
    print(f"exported {len(body)} relations")


if __name__ == "__main__":
    main()
