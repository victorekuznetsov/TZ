#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Извлекает текст из документов базы знаний для полнотекстового поиска.

Поддержаны PDF (pypdf), DOCX/PPTX (python-docx/python-pptx) и легаси
.doc/.ppt (Office 97-2003) — свой разбор бинарного формата (OLE2 через
`olefile`, `pip install olefile`), без LibreOffice: headless-конвертация
не работает в этой песочнице («source file could not be loaded» даже на
заведомо исправном .docx — глубже, чем формат файла, см. quality.json,
диагностировано strace: процесс завершается кодом 0, файл открывается,
ошибка не системная, а внутренняя логика загрузчика). Извлечение по
структуре формата (не перебор printable-строк):
  .doc — FIB → таблица кусков (Clx/PlcPcd) → сами куски текста
         (CP1252-сжатые или несжатые UTF-16LE), по [MS-DOC].
  .ppt — поток «PowerPoint Document», плоский проход по записям в
         поисках TextCharsAtom/TextBytesAtom (0x0FA0/0x0FA8), по [MS-PPT].
Хватает для полнотекстового поиска (не для точного рендеринга — колонтитулы
и содержимое таблиц/сносок попадают в общий текст как есть, без разметки).

Каталоги PDF (сканы без текстового слоя) уже исключены на уровне
build_kb.py — сюда не попадают вовсе.

Текст на документ обрезается (по умолчанию 60 000 знаков — этого
хватает на первые ~30-40 страниц типового руководства) — иначе индекс
по 390-страничному тому раздувает файл без пользы для поиска.

Источник читается напрямую из git-объектов ветки rawdata (`git cat-file`)
— локальная копия 2,5 ГБ не требуется.

Выход: data/kb_text.json — {path: {text, chars, truncated}}

Запуск:
  python3 build/build_kb_text.py <repo_dir> <kb.json> <out_dir> [--limit N]
"""
import sys, os, json, subprocess, io, re, signal, struct

MAX_CHARS = 60000
MAX_PAGES = 80  # страховка: не листать сотни страниц ради текста, если он редкий
DOC_TIMEOUT_SEC = 20  # pypdf иногда патологически виснет на одном повреждённом PDF


class TimeoutError_(Exception):
    pass


def _alarm(signum, frame):
    raise TimeoutError_()
SKIP_CLASSES = {"Чертежи CAD", "Фото", "Справочные таблицы"}
SUPPORTED_EXT = {".pdf", ".docx", ".pptx", ".doc", ".ppt"}


def git_show(repo_dir, path):
    r = subprocess.run(["git", "-C", repo_dir, "cat-file", "-p", f"origin/rawdata:{path}"],
                        capture_output=True)
    if r.returncode != 0:
        return None
    return r.stdout


def extract_pdf(data):
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    out = []
    total = 0
    for i, page in enumerate(r.pages):
        if i >= MAX_PAGES:
            break
        t = page.extract_text() or ""
        out.append(t)
        total += len(t)
        if total > MAX_CHARS:
            break
    return "\n".join(out)


def extract_docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_pptx(data):
    from pptx import Presentation
    p = Presentation(io.BytesIO(data))
    parts = []
    for slide in p.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)


def extract_doc(data):
    """Легаси .doc (Word 97-2003): FIB -> Clx/PlcPcd (таблица кусков) -> текст."""
    import olefile
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        if not ole.exists("WordDocument"):
            raise ValueError("нет потока WordDocument")
        wd = ole.openstream("WordDocument").read()
        if len(wd) < 32 or struct.unpack("<H", wd[0:2])[0] != 0xA5EC:
            raise ValueError("не похоже на WordDocument (wIdent)")
        flags = struct.unpack("<H", wd[10:12])[0]
        fWhichTblStream = (flags >> 9) & 1
        table_name = "1Table" if fWhichTblStream else "0Table"
        if not ole.exists(table_name):
            raise ValueError(f"нет потока {table_name}")
        table = ole.openstream(table_name).read()

        base = 32
        csw = struct.unpack("<H", wd[base:base + 2])[0]
        off = base + 2 + csw * 2
        cslw = struct.unpack("<H", wd[off:off + 2])[0]
        off += 2 + cslw * 4
        cbRgFcLcb = struct.unpack("<H", wd[off:off + 2])[0]
        off += 2
        fib_rg_fc_lcb = wd[off:off + cbRgFcLcb * 8]

        def fc_lcb(idx):
            o = idx * 8
            fc, lcb = struct.unpack("<II", fib_rg_fc_lcb[o:o + 8])
            return fc, lcb

        fcClx, lcbClx = fc_lcb(33)
        if lcbClx == 0:
            raise ValueError("lcbClx=0 — clx не найден")
        clx = table[fcClx:fcClx + lcbClx]

        i = 0
        pcdt = None
        while i < len(clx):
            clxt = clx[i]
            if clxt == 0x01:
                i += 1
                cb = struct.unpack("<H", clx[i:i + 2])[0]
                i += 2 + cb
            elif clxt == 0x02:
                i += 1
                lcb = struct.unpack("<I", clx[i:i + 4])[0]
                i += 4
                pcdt = clx[i:i + lcb]
                break
            else:
                break
        if pcdt is None:
            raise ValueError("Pcdt не найден в Clx")

        n = (len(pcdt) - 4) // 12
        cps = struct.unpack("<%dI" % (n + 1), pcdt[:4 * (n + 1)])
        pcd_start = 4 * (n + 1)
        out = []
        total = 0
        for k in range(n):
            pcd = pcdt[pcd_start + k * 8: pcd_start + (k + 1) * 8]
            fc_raw = struct.unpack("<I", pcd[2:6])[0]
            fCompressed = (fc_raw >> 30) & 1
            fc = fc_raw & 0x3FFFFFFF
            cp_start, cp_end = cps[k], cps[k + 1]
            n_chars = cp_end - cp_start
            if fCompressed:
                fc_byte = fc // 2
                raw = wd[fc_byte:fc_byte + n_chars]
                text = raw.decode("cp1252", "replace")
            else:
                raw = wd[fc:fc + n_chars * 2]
                text = raw.decode("utf-16-le", "replace")
            out.append(text)
            total += len(text)
            if total > MAX_CHARS:
                break
        full = "".join(out)
        return "".join(c if (c.isprintable() or c in "\n\t") else " " for c in full)
    finally:
        ole.close()


def extract_ppt(data):
    """Легаси .ppt (PowerPoint 97-2003): плоский проход по TextCharsAtom/TextBytesAtom."""
    import olefile
    TEXT_CHARS = 0x0FA0
    TEXT_BYTES = 0x0FA8
    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        if not ole.exists("PowerPoint Document"):
            raise ValueError("нет потока PowerPoint Document")
        d = ole.openstream("PowerPoint Document").read()
        out = []
        i = 0
        n = len(d)
        total = 0
        while i + 8 <= n:
            rec_ver_inst, rec_type, rec_len = struct.unpack("<HHI", d[i:i + 8])
            rec_ver = rec_ver_inst & 0x0F
            if rec_type in (TEXT_CHARS, TEXT_BYTES) and rec_ver == 0 \
                    and 0 < rec_len <= n - i - 8:
                body = d[i + 8:i + 8 + rec_len]
                text = body.decode("utf-16-le", "replace") if rec_type == TEXT_CHARS \
                    else body.decode("cp1252", "replace")
                if text.strip():
                    out.append(text)
                    total += len(text)
                i += 8 + rec_len
                if total > MAX_CHARS:
                    break
            else:
                i += 2
        full = "\n".join(out)
        return "".join(c if (c.isprintable() or c in "\n\t") else " " for c in full)
    finally:
        ole.close()


def clean(text):
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def main():
    args = sys.argv[1:]
    limit = None
    if "--limit" in args:
        i = args.index("--limit")
        limit = int(args[i + 1])
        args = args[:i] + args[i + 2:]
    repo_dir, kb_path, out_dir = args[:3]
    os.makedirs(out_dir, exist_ok=True)

    kb = json.load(open(kb_path, encoding="utf-8"))
    docs = [d for d in kb["docs"] if d["class"] not in SKIP_CLASSES
            and os.path.splitext(d["name"])[1].lower() in SUPPORTED_EXT]
    if limit:
        docs = docs[:limit]
    print(f"кандидатов на извлечение: {len(docs)}", file=sys.stderr)

    result = {}
    ok = fail = timedout = 0
    have_alarm = hasattr(signal, "SIGALRM")
    if have_alarm:
        signal.signal(signal.SIGALRM, _alarm)
    for n, d in enumerate(docs):
        ext = os.path.splitext(d["name"])[1].lower()
        data = git_show(repo_dir, d["path"])
        if data is None:
            fail += 1
            continue
        if have_alarm:
            signal.alarm(DOC_TIMEOUT_SEC)
        try:
            if ext == ".pdf":
                text = extract_pdf(data)
            elif ext == ".docx":
                text = extract_docx(data)
            elif ext == ".pptx":
                text = extract_pptx(data)
            elif ext == ".doc":
                text = extract_doc(data)
            elif ext == ".ppt":
                text = extract_ppt(data)
            else:
                continue
            text = clean(text)
            if not text:
                continue
            truncated = len(text) > MAX_CHARS
            result[d["path"]] = {"text": text[:MAX_CHARS], "chars": len(text), "truncated": truncated}
            ok += 1
        except TimeoutError_:
            timedout += 1
            print(f"  ТАЙМАУТ ({DOC_TIMEOUT_SEC}с): {d['path']}", file=sys.stderr)
        except Exception:
            fail += 1
        finally:
            if have_alarm:
                signal.alarm(0)
        if (n + 1) % 25 == 0:
            print(f"  {n+1}/{len(docs)}  ok={ok} fail={fail} timeout={timedout}", file=sys.stderr)

    out = {
        "meta": {
            "docsIndexed": ok, "docsFailed": fail, "docsTimedOut": timedout, "candidatesTotal": len(docs),
            "maxCharsPerDoc": MAX_CHARS,
        },
        "docs": result,
    }
    with open(os.path.join(out_dir, "kb_text.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print(f"\nпроиндексировано: {ok}  не удалось: {fail}  таймаут: {timedout}  из {len(docs)}")


if __name__ == "__main__":
    main()
