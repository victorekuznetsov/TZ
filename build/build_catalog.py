#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Строит каталог запчастей WK: прайс ДП (основной) + прайс УСО (сверка цены)
+ привязка к дереву узлов (tree.json) + привязка к коду ЕКМТР (mtr.json).

Цена в юанях используется ТОЛЬКО здесь, в каталоге — нигде больше в портале.

Выход:
  data/catalog.json   — 3042 позиции прайса со всеми связками
  data/ekmtr_wk.json  — срез НСИ: 1333 кода ППЗ 3.1.2.7 «Запчасти к WK»
                         (код, наименование, каталожный №, изготовитель) —
                         используется в кодификации, запасах и закупках,
                         без необходимости грузить весь mtr.json (20 МБ)

Запуск:
  python3 build/build_catalog.py <price_dp.xlsx> <price_uso.xlsx> \
      <mtr.json> <tree.json> <out_dir>
"""
import sys, os, re, json
import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from ekmtr_match import norm, extract_wk_group, build_index, match


def load_price_dp(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r[4]:
            continue
        out.append({
            "art": str(r[4]).strip(),
            "artNew": str(r[5]).strip() if r[5] else None,
            "model": r[1], "type": r[2], "zhClass": r[3],
            "nameRu": r[6], "nameZh": r[7],
            "resource": r[8], "tnved": r[9], "duty": r[10],
            "priceDP": r[11] if isinstance(r[11], (int, float)) else None,
            "priceDP2": r[12] if isinstance(r[12], (int, float)) else None,
            "note": r[13],
        })
    wb.close()
    return out


def load_price_uso(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        if not r[3]:
            continue
        art = str(r[3]).strip()
        out[art] = {
            "priceUSO": r[8] if isinstance(r[8], (int, float)) else None,
            "resourceUSO": r[6],
        }
    wb.close()
    return out


def build_tree_index(tree):
    """num (без учёта книги) -> список узлов дерева, где он встречается."""
    idx = {}
    for n in tree["nodes"]:
        idx.setdefault(n["num"], []).append(n)
    norm_idx = {}
    for num, entries in idx.items():
        norm_idx.setdefault(norm(num), []).extend(entries)
    return idx, norm_idx


def match_tree(art, art2, idx, norm_idx):
    for a in (art, art2):
        if not a:
            continue
        if a.upper() in idx:
            return idx[a.upper()], "exact"
    for a in (art, art2):
        if not a:
            continue
        na = norm(a)
        if na in norm_idx:
            return norm_idx[na], "norm"
    # родительский узел: ведомость взаимозаменяемости (2299 номеров) — подмножество
    # полного состава машины (3042 в прайсе), у не попавших туда деталей часто есть
    # узел-родитель на уровень выше (K1601.30.03.34 -> узел K1601.30.03 «Ковш»).
    # Сам номер детали в дереве не значится — это НЕ то же самое, что точное
    # совпадение, поэтому метод помечается отдельно: «через родителя».
    for a in (art, art2):
        if not a or not re.match(r"^K\d{4}", a, re.I):
            continue
        parts = a.upper().split(".")
        for k in range(len(parts) - 1, 0, -1):
            prefix = ".".join(parts[:k])
            if prefix in idx:
                return idx[prefix], "parent"
    return [], None


def main():
    if len(sys.argv) < 6:
        print("usage: build_catalog.py <price_dp.xlsx> <price_uso.xlsx> "
              "<mtr.json> <tree.json> <out_dir>", file=sys.stderr)
        sys.exit(1)
    dp_path, uso_path, mtr_path, tree_path, out_dir = sys.argv[1:6]
    os.makedirs(out_dir, exist_ok=True)

    dp = load_price_dp(dp_path)
    uso = load_price_uso(uso_path)
    tree = json.load(open(tree_path, encoding="utf-8"))
    tree_idx, tree_norm_idx = build_tree_index(tree)

    mtr = json.load(open(mtr_path, encoding="utf-8"))
    wk_group = extract_wk_group(mtr)
    by_cat, by_name_sub = build_index(wk_group)
    del mtr  # 20 МБ, больше не нужен

    dup = {}
    for row in dp:
        dup.setdefault(row["art"], []).append(row)
    seen = set()

    items = []
    n_ekmtr = n_tree = n_uso = n_ambiguous = 0
    for row in dp:
        art, art2 = row["art"], row["artNew"]
        uso_row = uso.get(art) or (uso.get(art2) if art2 else None)
        priceDiff = None
        if uso_row and uso_row["priceUSO"] is not None and row["priceDP"] is not None:
            d = row["priceDP"] - uso_row["priceUSO"]
            if abs(d) > 0.01:
                priceDiff = round(d, 2)

        nodes, tmethod = match_tree(art, art2, tree_idx, tree_norm_idx)
        tree_refs = [{"book": n["book"], "num": n["num"], "mech": n["mech"],
                      "mechNum": n["mechNum"], "qty": n["qty"], "top": n["top"]}
                     for n in nodes]

        code, method, ambiguous = match(art, art2, by_cat, by_name_sub)

        if uso_row:
            n_uso += 1
        if tree_refs:
            n_tree += 1
        if code:
            n_ekmtr += 1
        if ambiguous:
            n_ambiguous += 1

        items.append({
            "art": art, "artNew": art2,
            "model": row["model"], "type": row["type"], "zhClass": row["zhClass"],
            "nameRu": row["nameRu"], "nameZh": row["nameZh"],
            "resource": row["resource"], "tnved": row["tnved"], "duty": row["duty"],
            "priceCNY": row["priceDP"],
            "priceUsoCNY": uso_row["priceUSO"] if uso_row else None,
            "priceDiffCNY": priceDiff,
            "note": row["note"],
            "tree": tree_refs, "treeMethod": tmethod,
            "ekmtr": code, "ekmtrMethod": method, "ekmtrAmbiguous": ambiguous,
            "dupArticle": len(dup[art]) > 1,
        })

    priceOnlyUso = sorted(set(uso) - {r["art"] for r in dp} - {r["artNew"] for r in dp if r["artNew"]})

    catalog = {
        "meta": {
            "srcDP": os.path.basename(dp_path), "srcUSO": os.path.basename(uso_path),
            "currency": "CNY", "note": "Цена в юанях используется только в каталоге",
            "items": len(items),
            "matchedTree": n_tree, "matchedEkmtr": n_ekmtr, "matchedUso": n_uso,
            "ambiguousEkmtr": n_ambiguous,
            "priceOnlyInUso": len(priceOnlyUso),
        },
        "items": items,
        "priceOnlyInUso": priceOnlyUso,
    }
    with open(os.path.join(out_dir, "catalog.json"), "w", encoding="utf-8") as f:
        json.dump(catalog, f, ensure_ascii=False, separators=(",", ":"))

    ekmtr_wk = {
        "meta": {"src": "mtr.json ППЗ 3.1.2.7 «Запчасти к экскаваторам WK»",
                  "codes": len(wk_group)},
        "items": wk_group,
    }
    with open(os.path.join(out_dir, "ekmtr_wk.json"), "w", encoding="utf-8") as f:
        json.dump(ekmtr_wk, f, ensure_ascii=False, separators=(",", ":"))

    print(f"позиций прайса: {len(items)}")
    print(f"  привязано к дереву узлов: {n_tree} ({100*n_tree/len(items):.1f}%)")
    print(f"  привязано к коду ЕКМТР:   {n_ekmtr} ({100*n_ekmtr/len(items):.1f}%)  из них неоднозначно: {n_ambiguous}")
    print(f"  сверено с прайсом УСО:    {n_uso} ({100*n_uso/len(items):.1f}%)")
    print(f"  только в УСО (нет в ДП):  {len(priceOnlyUso)}")
    print(f"код ЕКМТР группы WK: {len(wk_group)}")


if __name__ == "__main__":
    main()
